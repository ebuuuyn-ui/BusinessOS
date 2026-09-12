"""Isolated regression checks; never opens the installed application's data."""
import os
import tempfile
import unittest
from decimal import Decimal
from io import BytesIO

_test_data = tempfile.TemporaryDirectory(prefix="businessos-report-test-")
os.environ["BUSINESSOS_DATA_DIR"] = _test_data.name
os.environ["DATABASE_URL"] = "sqlite:///:memory:"

from app import app, db, Customer, Order, OrderItem, AccountTransaction, Product, Invoice, InvoiceItem, normalize_search_text
from customer_movement_report import build_report
from openpyxl import load_workbook
from desktop import EXPORT_PATH


class ReportTests(unittest.TestCase):
    def setUp(self):
        app.config["TESTING"] = True
        self.context = app.app_context()
        self.context.push()
        db.drop_all()
        db.create_all()
        self.customer = Customer(name="İNCİ Müşteri", code="C01")
        db.session.add_all([self.customer, Customer(name="Hareketsiz")])
        db.session.flush()

    def tearDown(self):
        db.session.remove()
        self.context.pop()

    def order(self, no, kind, status, price, vat=0, discount=0):
        order = Order(customer=self.customer, order_no=no, order_type=kind, status=status)
        order.items.append(OrderItem(product_name="=Yedek parça", quantity=2, unit_price=Decimal(price), vat_rate=vat, discount_rate=discount, variant="Siyah", detail_2="Montaj dahil", detail_3="Özel", note="Test ayrıntı"))
        db.session.add(order)
        return order

    def movement(self, debit=0, credit=0):
        db.session.add(AccountTransaction(customer=self.customer, transaction_type="Ödeme" if debit else "Tahsilat", debit=debit, credit=credit, description="Test", payment_method="Banka"))

    def report(self, q=""):
        db.session.commit()
        return build_report(Customer, Order, AccountTransaction, normalize_search_text, q)

    def test_signed_balances_statuses_discount_and_no_invoice_double_count(self):
        delivered = self.order("SS1", "Satış", "Teslim Edildi", "100", vat=20, discount=10)  # 216
        self.order("SA1", "Satın Alma", "Teslim Edildi", "20")  # 40
        self.order("SS2", "Satış", "Sevk Edildi", "50")  # 100 delivered
        self.order("SA3", "Satın Alma", "Sevk Edildi", "15")  # 30 delivered
        self.order("SA2", "Satın Alma", "Bekliyor", "10")  # 20 pending
        self.order("CANCEL", "Satış", "İptal Edildi", "99999")
        self.movement(credit=100)
        self.movement(debit=10)
        product = Product(name="Stok")
        db.session.add(product)
        db.session.flush()
        invoice = Invoice(customer=self.customer, order=delivered, invoice_type="Satış", invoice_no="F1")
        invoice.items.append(InvoiceItem(product=product, product_name="Stok", quantity=1, unit_price=216, vat_rate=0))
        db.session.add(invoice)
        rows, totals = self.report("inci")
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["remaining"], 156)
        self.assertEqual(rows[0]["projected"], 136)
        self.assertEqual(totals["delivered_sale"], 316)
        self.assertEqual(totals["delivered_purchase"], 70)
        self.assertEqual(totals["pending_sale"], 0)
        self.assertEqual(len(rows[0]["delivered"]), 4)
        self.assertEqual(len(rows[0]["pending"]), 1)
        self.assertEqual(self.customer.balance, 126)  # Invoice ledger stays independent.

    def test_advance_and_payment_only_customer_are_not_hidden(self):
        self.order("SS1", "Satış", "Bekliyor", "50")
        self.movement(credit=150)
        rows, _ = self.report()
        self.assertEqual(rows[0]["remaining"], -150)
        self.assertEqual(rows[0]["projected"], -50)
        db.session.query(OrderItem).delete()
        db.session.query(Order).delete()
        rows, _ = self.report()
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["projected"], -150)

    def test_routes_excel_detail_filter_and_native_download(self):
        self.order("SS1", "Satış", "Teslim Edildi", "100")
        self.order("SS2", "Satış", "Bekliyor", "40")
        self.order("SS3", "Satış", "Sevk Edildi", "25")
        self.movement(credit=50)
        self.report()
        client = app.test_client()
        path = "/musteriler/toplam-hareket-bakiyeleri"
        response = client.get(path)
        self.assertEqual(response.status_code, 200)
        self.assertIn(b"SS1", response.data)
        self.assertIn(b"SS2", response.data)
        self.assertIn(b"data-workspace-new-tab", response.data)
        self.assertTrue(EXPORT_PATH.fullmatch(path + "/excel"))
        response = client.get(path + "/excel?q=inci")
        self.assertEqual(response.status_code, 200)
        book = load_workbook(BytesIO(response.data))
        self.assertEqual(len(book.sheetnames), 4)
        self.assertEqual(book["Cari Özet"]["G5"].value, 200)
        self.assertEqual(book["Cari Özet"]["J5"].value, 280)
        self.assertEqual(book["Teslim Edilen Kalemler"]["D6"].value, "SS3")
        self.assertEqual(book["Teslim Edilen Kalemler"]["G6"].value, "Sevk Edildi")
        self.assertEqual(book["Teslim Bekleyen Kalemler"].max_row, 5)
        for title in ["Teslim Edilen Kalemler", "Teslim Bekleyen Kalemler"]:
            sheet = book[title]
            self.assertEqual(sheet["H5"].value, "=Yedek parça")
            self.assertEqual(sheet["H5"].data_type, "s")
            self.assertEqual(sheet["J5"].value, "Montaj dahil")
        empty = load_workbook(BytesIO(client.get(path + "/excel?q=no-match").data))
        self.assertEqual(empty["Cari Özet"].max_row, 4)


if __name__ == "__main__":
    unittest.main()
