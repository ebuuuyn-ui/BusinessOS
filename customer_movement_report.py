"""Read-only, order-based customer exposure report (not the invoice ledger)."""
from collections import defaultdict
from datetime import date
from decimal import Decimal
from io import BytesIO

from flask import render_template, request, send_file
from sqlalchemy.orm import selectinload

TITLE = "Cari Toplam Hareket Bakiyeleri"
NOTE = ("Sipariş bazlı rapor: teslim edilen satış − teslim edilen satın alma + cari borç hareketleri − cari alacak hareketleri = kalan bakiye. "
        "Sevk Edildi ve Teslim Edildi durumları teslim edilenler grubuna dahildir. "
        "Bekleyen satışlar eklenir, bekleyen satın almalar düşülür. Pozitif bakiye carinin borcu, negatif bakiye carinin alacağıdır. "
        "Tutarlar iskonto sonrası KDV dahildir. Faturalar ayrıca eklenmez; bu rapor faturaya dayalı cari hesap ekstresi değildir.")
ZERO = Decimal("0")
KEYS = ("delivered_sale", "delivered_purchase", "debit", "credit", "remaining", "pending_sale", "pending_purchase", "projected")


def build_report(Customer, Order, Transaction, normalize, query=""):
    orders = defaultdict(list)
    movements = defaultdict(list)
    for order in Order.query.options(selectinload(Order.items)).filter(Order.status != "İptal Edildi").order_by(Order.order_date, Order.id).all():
        orders[order.customer_id].append(order)
    for movement in Transaction.query.order_by(Transaction.transaction_date, Transaction.id).all():
        movements[movement.customer_id].append(movement)
    rows = []
    terms = normalize(query).split()
    for customer in Customer.query.order_by(Customer.name).all():
        if customer.id not in orders and customer.id not in movements:
            continue
        if not all(term in normalize(f"{customer.name} {customer.code or ''}") for term in terms):
            continue
        row = dict.fromkeys(KEYS, ZERO)
        row.update(customer=customer, delivered=[], pending=[], movements=movements[customer.id])
        for order in orders[customer.id]:
            group = "delivered" if order.status in {"Teslim Edildi", "Sevk Edildi"} else "pending"
            row[group].append(order)
            kind = "sale" if order.order_type == "Satış" else "purchase"
            row[f"{group}_{kind}"] += order.total_amount
        row["debit"] = sum((m.debit or ZERO for m in row["movements"]), ZERO)
        row["credit"] = sum((m.credit or ZERO for m in row["movements"]), ZERO)
        row["remaining"] = row["delivered_sale"] - row["delivered_purchase"] + row["debit"] - row["credit"]
        row["projected"] = row["remaining"] + row["pending_sale"] - row["pending_purchase"]
        rows.append(row)
    totals = {key: sum((row[key] for row in rows), ZERO) for key in KEYS}
    return rows, totals


def build_workbook(rows, query):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    book = Workbook()
    book.remove(book.active)
    money_format = '#,##0.00;[Red]-#,##0.00'

    def sheet(name, headers, data, money_columns, widths):
        ws = book.create_sheet(name)
        ws.append([TITLE])
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
        ws.append([NOTE])
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(headers))
        ws.cell(2, 1).alignment = Alignment(wrap_text=True, vertical="center")
        ws.row_dimensions[2].height = 50
        ws.append([f"Rapor tarihi: {date.today():%d.%m.%Y} · Cari araması: {query or 'Tüm cariler'} · Tüm kayıt tarihleri"])
        ws.append(headers)
        for values in data:
            ws.append([float(value) if isinstance(value, Decimal) else value for value in values])
        for row in ws.iter_rows():
            for cell in row:
                # Customer-entered labels must stay text, never Excel formulas.
                if cell.data_type == "f":
                    cell.data_type = "s"
                cell.font = Font(name="Calibri", size=11)
                if cell.row >= 5:
                    cell.alignment = Alignment(vertical="top", wrap_text=cell.column not in money_columns)
                    if cell.column in money_columns:
                        cell.number_format = money_format
                    if cell.row % 2:
                        cell.fill = PatternFill("solid", fgColor="F3F6FA")
                    if isinstance(cell.value, date):
                        cell.number_format = "dd.mm.yyyy"
        ws.cell(1, 1).font = Font(size=17, bold=True, color="14213D")
        for cell in ws[4]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="24456A")
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.row_dimensions[4].height = 34
        for i, width in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = width
        ws.freeze_panes = "C5"
        ws.auto_filter.ref = f"A4:{get_column_letter(len(headers))}{max(4, ws.max_row)}"
        ws.sheet_view.showGridLines = False
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.print_title_rows = "1:4"
        return ws

    summary = sheet("Cari Özet", ["Cari Kodu", "Cari Adı", "Teslim Edilen Satış", "Teslim Edilen Satın Alma", "Ödeme / Diğer Borç", "Tahsilat / Diğer Alacak", "Kalan Bakiye", "Bekleyen Satış", "Bekleyen Satın Alma", "Toplam Beklenen Bakiye"],
          [[r["customer"].code or "", r["customer"].name] + [r[k] for k in KEYS] for r in rows], set(range(3, 11)), [18, 42] + [22] * 8)
    if rows:
        summary.append(["TOPLAM", ""] + [float(sum((r[k] for r in rows), ZERO)) for k in KEYS])
        for cell in summary[summary.max_row]:
            cell.font = Font(bold=True, color="14213D")
            cell.fill = PatternFill("solid", fgColor="E8EEF9")
            if cell.column >= 3:
                cell.number_format = money_format
    headers = ["Cari Kodu", "Cari Adı", "Tür", "Sipariş No", "Sipariş Tarihi", "Teslim Tarihi", "Durum", "Ürün / Hizmet", "Ayrıntı 1", "Ayrıntı 2", "Ayrıntı 3", "Adet", "Birim", "Birim Fiyat", "Fiyat KDV", "İskonto %", "KDV %", "KDV Hariç", "KDV", "KDV Dahil", "Kalem Notu", "Sipariş Notu"]
    for group, name in [("delivered", "Teslim Edilen Kalemler"), ("pending", "Teslim Bekleyen Kalemler")]:
        data = []
        for r in rows:
            c = r["customer"]
            for o in r[group]:
                for i in o.items:
                    data.append([c.code or "", c.name, o.order_type, o.order_no, o.order_date, o.delivery_date, o.status, i.product_name, i.variant, i.detail_2, i.detail_3, i.quantity, i.unit, i.unit_price, "Dahil" if i.vat_included else "Hariç", i.discount_rate, i.vat_rate, i.net_amount, i.vat_amount, i.total_amount, i.note, o.notes])
        sheet(name, headers, data, {14, 18, 19, 20}, [18, 40, 16, 22, 16, 16, 22, 42, 24, 24, 24, 10, 12, 18, 12, 12, 12, 18, 18, 18, 35, 35])
    sheet("Tahsilat ve Ödemeler", ["Cari Kodu", "Cari Adı", "Tarih", "Hareket Türü", "Referans", "Açıklama", "Ödeme Şekli", "Borç (+)", "Alacak (−)"],
          [[r["customer"].code or "", r["customer"].name, m.transaction_date, m.transaction_type, m.reference_no, m.description, m.payment_method, m.debit or ZERO, m.credit or ZERO] for r in rows for m in r["movements"]], {8, 9}, [18, 42, 16, 24, 24, 48, 20, 20, 20])
    output = BytesIO()
    book.save(output)
    output.seek(0)
    return output


def register_report(app, Customer, Order, Transaction, normalize):
    @app.get("/musteriler/toplam-hareket-bakiyeleri")
    def customer_total_movements():
        query = request.args.get("q", "").strip()
        rows, totals = build_report(Customer, Order, Transaction, normalize, query)
        return render_template("customer_movement_report.html", rows=rows, totals=totals, query=query, report_note=NOTE)

    @app.get("/musteriler/toplam-hareket-bakiyeleri/excel")
    def export_customer_total_movements():
        query = request.args.get("q", "").strip()
        rows, _ = build_report(Customer, Order, Transaction, normalize, query)
        return send_file(build_workbook(rows, query), as_attachment=True,
                         download_name=f"Cari-Toplam-Hareket-Bakiyeleri-{date.today().isoformat()}.xlsx",
                         mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
