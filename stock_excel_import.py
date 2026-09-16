"""Previewed, atomic stock-card replacement with protection for used products."""
import csv
import hashlib
import json
import unicodedata
from collections import Counter, defaultdict
from datetime import date, datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from io import BytesIO, StringIO
from uuid import uuid4

from flask import abort, redirect, render_template, request, send_file, url_for
from itsdangerous import URLSafeTimedSerializer, BadSignature
from openpyxl import load_workbook
from sqlalchemy import func, select


def norm(value):
    return ''.join(c for c in unicodedata.normalize('NFKD', str(value or '').casefold().replace('ı','i')) if not unicodedata.combining(c)).strip()


def read_products(content, skip_missing=False):
    workbook=load_workbook(BytesIO(content), read_only=True, data_only=True)
    aliases={'code':{'urun kodu','stok kodu','kart kodu'},'name':{'stok adi'},
             'unit_price':{'liste fiyati','liste fiyat (tl)','liste fiyati (tl)'},
             'purchase_price':{'abika satis fiyati'}}
    try:
        candidates=[]
        for sheet in workbook:
            for number,row in enumerate(sheet.iter_rows(values_only=True),1):
                if number>20:break
                headers={norm(v):i for i,v in enumerate(row) if v is not None}
                columns={key:next((headers[x] for x in names if x in headers),None) for key,names in aliases.items()}
                if all(v is not None for v in columns.values()):candidates.append((sheet,number,columns));break
        if len(candidates)!=1:raise ValueError('Stok Adı, Ürün Kodu, Liste Fiyatı ve ABİKA Satış Fiyatı başlıklarını içeren tek bir sayfa gerekli.')
        sheet,header,columns=candidates[0]
        records=[];skipped=[];codes=set()
        for number,row in enumerate(sheet.iter_rows(min_row=header+1,values_only=True),header+1):
            if all(v is None for v in row):continue
            values={k:row[i] if i<len(row) else None for k,i in columns.items()}
            code=str(values['code'] or '').strip();name=str(values['name'] or '').strip()
            if not code or not name or len(code)>80 or len(name)>160:raise ValueError(f'Satır {number}: ürün kodu veya stok adı eksik/geçersiz.')
            if code in codes:raise ValueError(f'Tekrarlanan ürün kodu: {code}')
            codes.add(code)
            if any(values[k] is None or str(values[k]).strip()=='' for k in ('unit_price','purchase_price')):
                if not skip_missing:raise ValueError(f'{code}: alış veya satış fiyatı boş. Eksik fiyatlı satırları atlama seçeneği kullanılabilir.')
                skipped.append({'code':code,'name':name,'row':number,'reason':'Eksik fiyat'});continue
            record={'code':code,'name':name,'row':number}
            for field in ('unit_price','purchase_price'):
                try:
                    raw=values[field]
                    if isinstance(raw,bool):raise ValueError()
                    value=Decimal(str(raw))
                    if not value.is_finite() or value<0 or value>=Decimal('10000000000'):raise ValueError()
                    record[field]=str(value.quantize(Decimal('.01'),rounding=ROUND_HALF_UP))
                except (InvalidOperation,ValueError):raise ValueError(f'{code}: geçersiz fiyat ({field}).')
            records.append(record)
        if not records:raise ValueError('Aktarılabilir ürün bulunamadı.')
        return records,skipped,sorted(codes)
    finally:workbook.close()


def snapshots(products):
    return [{c.name:(str(getattr(p,c.name)) if isinstance(getattr(p,c.name),(Decimal,date,datetime)) else getattr(p,c.name)) for c in p.__table__.columns} for p in products]


def protection(db,products):
    reasons=defaultdict(dict)
    labels={'order_item':'Sipariş','invoice_item':'Fatura','stock_movement':'Stok hareketi','ozon_sale':'Ozon satışı'}
    for table in db.metadata.tables.values():
        for column in table.columns:
            if any(f.target_fullname=='product.id' for f in column.foreign_keys):
                for pid,count in db.session.execute(select(column,func.count()).where(column.isnot(None)).group_by(column)):
                    reasons[pid][labels.get(table.name,table.name)]=count
                # Legacy lines without an FK may still identify a product by name.
                if 'product_name' in table.c:
                    names=defaultdict(int)
                    for name,count in db.session.execute(select(table.c.product_name,func.count()).where(column.is_(None)).group_by(table.c.product_name)):
                        if name:names[norm(name)]+=count
                    for product in products:
                        if names[norm(product.name)]:reasons[product.id]['Kodsuz '+labels.get(table.name,table.name)]=names[norm(product.name)]
    return dict(reasons)


def make_plan(products,reasons,records,source_codes,deactivate):
    by_code={p.code:p for p in products};incoming={r['code']:r for r in records};source=set(source_codes)
    changes=[];protected=[]
    for p in products:
        if p.id in reasons:
            protected.append({'id':p.id,'code':p.code or '', 'name':p.name,'reason':', '.join(f'{k}: {v}' for k,v in reasons[p.id].items())})
        elif p.code in incoming:
            r=incoming[p.code]
            if p.name!=r['name'] or p.unit_price!=Decimal(r['unit_price']) or p.purchase_price!=Decimal(r['purchase_price']) or not p.active:
                changes.append(dict(r,id=p.id,action='Güncelle'))
        elif deactivate and p.code not in source and p.active:
            changes.append({'id':p.id,'code':p.code or '', 'name':p.name,'action':'Pasife al'})
    for r in records:
        if r['code'] not in by_code:changes.append(dict(r,id=None,action='Yeni'))
    state={'products':snapshots(products),'protection':reasons}
    fingerprint=hashlib.sha256(json.dumps(state,sort_keys=True,ensure_ascii=False).encode()).hexdigest()
    return {'changes':changes,'protected':protected,'counts':dict(Counter(r['action'] for r in changes)), 'fingerprint':fingerprint,'before':state['products']}


def register_stock_excel_import(app,db,Product,StoredFile):
    signer=URLSafeTimedSerializer(app.secret_key,salt='stock-excel-v1')

    @app.route('/urunler/excel-guncelle',methods=['GET','POST'])
    def stock_excel_update():
        try:
            if request.method=='GET':return render_template('stock_excel_import.html')
            if request.form.get('action')=='apply':
                payload=signer.loads(request.form.get('plan',''),max_age=3600)
                key='stock-import/'+payload['nonce']+'.json'
                existing=StoredFile.query.filter_by(storage_key=key).first()
                if existing:return redirect(url_for('stock_excel_result',key=payload['nonce']))
                # PostgreSQL row locks also block concurrent new FK references.
                products=Product.query.order_by(Product.id).with_for_update().all()
                reasons=protection(db,products)
                plan=make_plan(products,reasons,payload['records'],payload['source_codes'],payload['deactivate'])
                if plan['fingerprint']!=payload['fingerprint']:
                    raise ValueError('Önizlemeden sonra stoklar veya hareketleri değişmiş. Dosyayı yeniden önizleyin; hiçbir kayıt değiştirilmedi.')
                by_id={p.id:p for p in products}
                for change in plan['changes']:
                    if change['action']=='Yeni':
                        p=Product(code=change['code'],name=change['name'],unit='Adet',active=True)
                        db.session.add(p)
                    else:p=by_id[change['id']]
                    if change['action']=='Pasife al':p.active=False
                    else:
                        p.name=change['name'];p.unit_price=Decimal(change['unit_price']);p.purchase_price=Decimal(change['purchase_price']);p.active=True
                db.session.flush()
                # Verify protected rows byte-for-byte at the column level before commit.
                original={p['id']:p for p in plan['before']}
                for row in snapshots([by_id[pid] for pid in reasons if pid in by_id]):
                    if row!=original[row['id']]:raise ValueError('Hareket görmüş ürün koruma kontrolü başarısız.')
                audit={'source':payload['source'],'source_sha256':payload['source_sha256'],'counts':plan['counts'],
                       'changes':plan['changes'],'protected':plan['protected'],'skipped':payload['skipped'],
                       'before':plan['before'],'after':snapshots(Product.query.order_by(Product.id).all())}
                content=json.dumps(audit,ensure_ascii=False,indent=2).encode()
                db.session.add(StoredFile(storage_key=key,content=content,mime_type='application/json',size_bytes=len(content),sha256=hashlib.sha256(content).hexdigest()))
                db.session.commit()
                return redirect(url_for('stock_excel_result',key=payload['nonce']))
            uploaded=request.files.get('file')
            if not uploaded or not uploaded.filename.lower().endswith('.xlsx'):raise ValueError('Bir .xlsx dosyası seçin.')
            content=uploaded.read(5*1024*1024+1)
            if len(content)>5*1024*1024:raise ValueError('Dosya en fazla 5 MB olabilir.')
            records,skipped,source_codes=read_products(content,request.form.get('skip_missing')=='on')
            products=Product.query.order_by(Product.id).all();reasons=protection(db,products)
            deactivate=request.form.get('deactivate')=='on'
            plan=make_plan(products,reasons,records,source_codes,deactivate)
            payload=dict(records=records,skipped=skipped,source_codes=source_codes,deactivate=deactivate,
                fingerprint=plan['fingerprint'],nonce=uuid4().hex,source=uploaded.filename,source_sha256=hashlib.sha256(content).hexdigest())
            return render_template('stock_excel_import.html',plan=plan,token=signer.dumps(payload),skipped=skipped,source_count=len(source_codes))
        except (ValueError,BadSignature) as exc:
            db.session.rollback()
            return render_template('stock_excel_import.html',error=str(exc) if isinstance(exc,ValueError) else 'Önizleme süresi dolmuş veya geçersiz. Yeniden önizleyin.'),400
        except Exception:
            db.session.rollback()
            app.logger.exception('Stock Excel import failed')
            return render_template('stock_excel_import.html',error='Aktarım tamamlanamadı. Değişiklikler geri alındı; dosyayı yeniden önizleyin.'),400

    @app.get('/urunler/excel-sonuc/<key>')
    def stock_excel_result(key):
        if len(key)!=32 or any(c not in '0123456789abcdef' for c in key):abort(404)
        record=StoredFile.query.filter_by(storage_key='stock-import/'+key+'.json').first_or_404()
        result=json.loads(record.content)
        if request.args.get('download')=='backup':
            return send_file(BytesIO(record.content),mimetype='application/json',as_attachment=True,download_name='stok-aktarim-yedegi-'+key+'.json')
        if request.args.get('download')=='protected':
            out=StringIO();writer=csv.writer(out,delimiter=';');writer.writerow(['Stok Kodu','Stok Adı','Koruma Nedeni'])
            for p in result['protected']:
                writer.writerow([("'"+v if v[:1] in '=+-@' else v) for v in [p['code'],p['name'],p['reason']]])
            return send_file(BytesIO(out.getvalue().encode('utf-8-sig')),mimetype='text/csv',as_attachment=True,download_name='korunan-stoklar.csv')
        return render_template('stock_excel_import.html',result=result,key=key)
