"""Read-only invoice register export."""
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


def build_invoice_register(invoices):
    book = Workbook()
    sheet = book.active
    sheet.title = 'Faturalar'
    sheet.append(['Fatura Tarihi', 'Vade Tarihi', 'Fatura No', 'Fatura Türü',
                  'Cari Kodu', 'Cari Adı', 'Vergi Dairesi', 'Vergi Numarası',
                  'Bağlı Sipariş', 'Kalem Sayısı', 'KDV Hariç (TL)', 'KDV (TL)',
                  'Genel Toplam (TL)', 'Not'])
    for invoice in invoices:
        customer = invoice.customer
        values = [invoice.invoice_date, invoice.due_date, invoice.invoice_no,
                  invoice.invoice_type, customer.code or '', customer.name,
                  customer.tax_office or '', customer.tax_number or '',
                  ', '.join(o.order_no for o in invoice.linked_orders) if hasattr(invoice,'linked_orders') else (invoice.order.order_no if invoice.order else ''), len(invoice.items),
                  invoice.net_amount, invoice.vat_amount, invoice.total_amount,
                  invoice.notes or '']
        sheet.append(values)
        for cell in sheet[sheet.max_row]:
            # Customer supplied text must remain text, never an Excel formula.
            if isinstance(cell.value, str): cell.data_type = 's'
            cell.alignment = Alignment(vertical='top', wrap_text=True)
        for column in (1, 2): sheet.cell(sheet.max_row, column).number_format = 'dd.mm.yyyy'
        for column in (11, 12, 13): sheet.cell(sheet.max_row, column).number_format = '#,##0.00'
    for cell in sheet[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='2563EB')
        cell.alignment = Alignment(wrap_text=True, vertical='center')
    sheet.row_dimensions[1].height = 32
    widths = [16, 16, 25, 20, 18, 48, 24, 20, 24, 15, 20, 18, 22, 50]
    for column, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_view.showGridLines = False
    output = BytesIO()
    book.save(output)
    output.seek(0)
    return output
