"""Read-only exports using the exact collection-tracking screen context."""
from maturity_labels import tracking_text
from datetime import date
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from xml.sax.saxutils import escape
import math

HEADERS = ['Cari', 'Fatura / Hareket', 'Belge Tarihi', 'Vade Tarihi', 'Tutar', 'Mahsup Edilen', 'Kalan', 'Durum']
STATES = {'open': 'Açık tahsilatlar', 'overdue': 'Geciken tahsilatlar', 'due_soon': '7 gün içinde vadeli', 'paid': 'Tahsil edilenler', 'all': 'Tümü', 'undated': 'Vadesi belirtilmemiş faturalar', 'other': 'Fatura dışı bakiyeler'}
NOTE = 'Tutarlar faturalar ve tüm cari hareketleriyle netleştirilir. Mahsuplar en eski fatura/hareketten düşülür; kayıtlı ödeme-fatura eşleştirmesi değildir. Vade yalnız faturadaki tarihtir; vade tarihi olmayan açık faturalar gecikmiş sayılmaz. Fatura dışı cari hareketlerine vade uygulanmaz.'


def filter_label(context):
    return f"Durum: {tracking_text(STATES[context['selected_state']], context)} | Arama: {context['query'] or 'Yok'}"


def rows(context):
    return [[i['customer'].name, i['reference'], i['document_date'], i['due_date'] or ('Vadesi belirtilmemiş' if i.get('invoice') is not None and i['remaining'] else '—'), i['amount'], i['collected'], i['remaining'], i['state_label']] for i in context['items']]


def totals(context):
    return [sum((i[key] for i in context['items']), Decimal('0')) for key in ('amount', 'collected', 'remaining')]


def export_excel(context):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    book = Workbook()
    ws = book.active
    ws.title = tracking_text('Tahsilat Takibi', context)
    summary = context['summary']
    top = [
        tracking_text('Tahsilat Takibi - Fatura Vadeleri ve Cari Bakiyeler', context),
        f"Rapor tarihi: {context['today']:%d.%m.%Y} | {filter_label(context)}",
        tracking_text(NOTE, context),
        'Genel Özet (filtrelerden bağımsız)',
    ]
    for text in top:
        ws.append([text])
        ws.merge_cells(start_row=ws.max_row, start_column=1, end_row=ws.max_row, end_column=8)
    ws.append([tracking_text('Geciken Tahsilatlar', context), float(summary['overdue_amount']), 'Kayıt sayısı', summary['overdue_count']])
    ws.append(['7 Gün İçinde Vadeli', float(summary['due_soon_amount']), 'Kayıt sayısı', summary['due_soon_count']])
    ws.append([tracking_text('Açık Tahsilat Toplamı', context), float(summary['open_amount'])])
    ws.append([f"Filtrelenen Liste - {len(context['items'])} kayıt"])
    ws.merge_cells('A8:H8')
    ws.append([tracking_text(h, context) for h in HEADERS])
    for row in rows(context):
        ws.append([float(v) if isinstance(v, Decimal) else v for v in row])
    last_data = ws.max_row
    if not context['items']:
        ws.append([tracking_text('Bu filtreye uygun fatura / cari hareketi bulunmuyor.', context)])
        ws.merge_cells(start_row=ws.max_row, start_column=1, end_row=ws.max_row, end_column=8)
    ws.append(['Filtrelenen Liste Toplamı', '', '', '', *map(float, totals(context)), ''])
    widths = [44, 23, 18, 18, 22, 22, 22, 26]
    for idx, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    for row in ws:
        for cell in row:
            if cell.data_type == 'f':
                cell.data_type = 's'
            cell.font = Font(name='Arial', size=10, color='172033')
            cell.alignment = Alignment(vertical='top', wrap_text=True)
            if isinstance(cell.value, date):
                cell.number_format = 'dd.mm.yyyy'
            if cell.row >= 10 and cell.column in (5, 6, 7) or cell.row in (5, 6, 7) and cell.column == 2:
                cell.number_format = '#,##0.00'
                cell.alignment = Alignment(horizontal='right', vertical='top')
            if cell.row in (4, 8, 9, ws.max_row):
                cell.fill = PatternFill('solid', fgColor='24456A' if cell.row == 9 else 'E9EFF8')
                cell.font = Font(name='Arial', bold=True, size=10, color='FFFFFF' if cell.row == 9 else '172033')
        lines = max((sum(max(1, math.ceil(len(part) / max(8, widths[c.column-1]-3))) for part in str(c.value or '').split('\n')) for c in row), default=1)
        ws.row_dimensions[row[0].row].height = max(28, min(409, lines*14))
    ws.row_dimensions[1].height = 30
    ws.row_dimensions[2].height = max(30, math.ceil(len(top[1])/140)*15)
    ws.row_dimensions[3].height = 32
    ws.cell(1, 1).font = Font(name='Arial', bold=True, size=16, color='172033')
    ws.freeze_panes = 'C10'
    ws.auto_filter.ref = f'A9:H{last_data}'
    ws.print_title_rows = '8:9'
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_view.showGridLines = False
    output = BytesIO()
    book.save(output)
    output.seek(0)
    return output


def export_pdf(context):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    regular = Path('/System/Library/Fonts/Supplemental/Arial.ttf')
    bold = regular.with_name('Arial Bold.ttf')
    if not regular.exists():
        regular = Path('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf')
        bold = regular.with_name('DejaVuSans-Bold.ttf')
    pdfmetrics.registerFont(TTFont('CollectionFont', str(regular)))
    pdfmetrics.registerFont(TTFont('CollectionBold', str(bold)))
    body = ParagraphStyle('body', fontName='CollectionFont', fontSize=8, leading=11, wordWrap='CJK')
    title = ParagraphStyle('title', parent=body, fontName='CollectionBold', fontSize=16, leading=21, spaceAfter=8)
    heading = ParagraphStyle('heading', parent=body, fontName='CollectionBold', fontSize=10, leading=14, spaceBefore=10, spaceAfter=6)
    header = ParagraphStyle('header', parent=body, fontName='CollectionBold', textColor=colors.white)
    right = ParagraphStyle('right', parent=body, alignment=2)
    para = lambda value, style=body: Paragraph(escape(str(value)).replace('\n', '<br/>'), style)
    money = lambda v: f'{v:,.2f}'.replace(',', 'X').replace('.', ',').replace('X', '.')
    width = landscape(A4)[0] - 48
    output = BytesIO()
    doc = SimpleDocTemplate(output, pagesize=landscape(A4), leftMargin=24, rightMargin=24, topMargin=24, bottomMargin=30, title=tracking_text('Tahsilat Takibi', context))
    story = [para(tracking_text('Tahsilat Takibi', context), title), para(f"Rapor tarihi: {context['today']:%d.%m.%Y} | {filter_label(context)}"), para(tracking_text(NOTE, context)), para('Genel Özet (filtrelerden bağımsız)', heading)]
    s = context['summary']
    story.append(para(f"Geciken: {money(s['overdue_amount'])} TL ({s['overdue_count']} kayıt) | 7 gün içinde vadeli: {money(s['due_soon_amount'])} TL ({s['due_soon_count']} kayıt) | Açık toplam: {money(s['open_amount'])} TL"))
    story.append(para(f"Filtrelenen Liste - {len(context['items'])} kayıt", heading))
    data = [[para(h, header) for h in [tracking_text(h, context) for h in HEADERS]]]
    for row in rows(context):
        data.append([para(money(value), right) if idx in (4,5,6) else para(value.strftime('%d.%m.%Y') if isinstance(value, date) else value) for idx, value in enumerate(row)])
    data.append([para('Liste Toplamı'), '', '', '', *[para(money(value), right) for value in totals(context)], ''])
    table = Table(data, colWidths=[width*p/100 for p in (24,14,10,10,11,11,11,9)], repeatRows=1, splitByRow=1, splitInRow=1)
    table.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#24456A')), ('BACKGROUND',(0,-1),(-1,-1),colors.HexColor('#E9EFF8')), ('VALIGN',(0,0),(-1,-1),'TOP'), ('LINEBELOW',(0,0),(-1,-1),.3,colors.HexColor('#D6DFEA')), ('ROWBACKGROUNDS',(0,1),(-1,-2),[colors.white,colors.HexColor('#F4F7FB')]), ('LEFTPADDING',(0,0),(-1,-1),5), ('RIGHTPADDING',(0,0),(-1,-1),5), ('TOPPADDING',(0,0),(-1,-1),6), ('BOTTOMPADDING',(0,0),(-1,-1),6)]))
    story.append(table)
    if not context['items']:
        story.extend([Spacer(1,8), para(tracking_text('Bu filtreye uygun fatura / cari hareketi bulunmuyor.', context))])
    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont('CollectionFont', 8)
        canvas.drawString(24, 14, tracking_text('Business OS | Tahsilat Takibi | Tutarlar TL', context))
        canvas.drawRightString(width+24, 14, f'Sayfa {doc.page}')
        canvas.restoreState()
    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    output.seek(0)
    return output
