import unittest
from io import BytesIO
from openpyxl import load_workbook
from test_seller_reports import SellerTests
from test_supplier_lead_times import LeadTests
from order_report_exports import export_report

class ExportTests(unittest.TestCase):
    def test_seller_numeric_values_and_untrusted_text(self):
        f=SellerTests();o=f.order();o.customer.name='=HYPERLINK("bad")'
        c=f.report([o]);wb=load_workbook(export_report('seller',c,'xlsx'))
        self.assertEqual(wb['Sipariş Ayrıntıları']['D5'].data_type,'s')
        self.assertEqual(wb['Sipariş Ayrıntıları']['I5'].value,120.3)
        self.assertEqual(wb['Satıcı Özeti']['F7'].value,120.3)
        self.assertTrue(export_report('seller',c,'pdf').read().startswith(b'%PDF'))
    def test_pending_lead_is_blank_and_empty_exports_work(self):
        f=LeadTests();c=f.report(f.order([(0,'Bekliyor')],status='Bekliyor'))
        wb=load_workbook(export_report('lead',c,'xlsx'));self.assertIsNone(wb['Sipariş Ayrıntıları']['F5'].value)
        self.assertEqual(wb['Sipariş Ayrıntıları']['H5'].value,'Devam Ediyor')
        for kind,context in [('lead',f.report()),('seller',SellerTests().report([]))]:
            self.assertTrue(export_report(kind,context,'pdf').read().startswith(b'%PDF'))
            self.assertEqual(len(load_workbook(export_report(kind,context,'xlsx')).sheetnames),2)
