"""Excel/PDF downloads from the exact filtered report context used on screen."""
from io import BytesIO
from datetime import date, datetime
from decimal import Decimal
from html import escape


def report_tables(kind,c):
    if kind=='seller':
        title='Satıcı Bazlı Sipariş Raporu'
        summary=[[g['name'],g['count'],g['quantity'],g['net'],g['vat'],g['total']] for g in c['seller_groups']]
        t=c['sales_totals'];summary.append(['Toplam',len(c['sales_rows']),t['quantity'],t['net'],t['vat'],t['total']])
        detail=[[r['order'].order_date,r['order'].order_no,r['seller'],r['order'].customer.name,r['order'].status,r['quantity'],r['net'],r['vat'],r['total']] for r in c['sales_rows']]
        note='İptaller hariçtir. Oluşturan kullanıcı kaydı olmayan eski siparişler Bekir’e dahildir. Tutarlar sipariş tutarıdır.'
        tables=[('Satıcı Özeti',['Satıcı','Sipariş Sayısı','Ürün Adedi','KDV Hariç (TL)','KDV (TL)','Toplam (TL)'],summary,[35,16,16,22,20,22]),('Sipariş Ayrıntıları',['Tarih','Sipariş No','Satıcı','Cari','Durum','Adet','KDV Hariç (TL)','KDV (TL)','Toplam (TL)'],detail,[15,21,18,45,22,12,20,18,20])]
    else:
        title='Tedarikçi Termin Raporu'
        summary=[[g['customer'].name,len(g['rows']),len(g['leads']),g['average'],g['pending'],g['missing'],g['cancelled']] for g in c['supplier_groups']]
        summary.append(['Genel Ortalama / Toplam',len(c['report_rows']),c['measured_count'],c['lead_average'],c['pending_count'],c['missing_count'],sum(g['cancelled'] for g in c['supplier_groups'])])
        detail=[[r['order'].order_no,r['order'].customer.name,r['created'],r['first'],r['first_status'],r['lead'],r['order'].status,r['state'],r['age']] for r in c['report_rows']]
        note='Takvim günü kullanılır; her sipariş eşit ağırlıktadır. Devam eden, iptal ve geçmişi eksik siparişler ortalamaya dahil değildir.'
        tables=[('Tedarikçi Özeti',['Tedarikçi','Sipariş','Ölçülen','Ortalama Termin (Gün)','Devam Eden','Geçmişi Eksik','İptal'],summary,[45,12,12,22,15,17,12]),('Sipariş Ayrıntıları',['Sipariş No','Tedarikçi','Sisteme Giriş','Geçerli Geçiş','Geçiş Durumu','Termin (Gün)','Güncel Durum','Ölçüm','Geçen Süre (Gün)'],detail,[21,40,23,23,22,16,22,28,18])]
    labels={'seller':'Satıcı','q':'Arama','customer_q':'Tedarikçi','start_date':'Başlangıç','end_date':'Bitiş'}
    filters=' · '.join(labels[k]+': '+str(v) for k,v in c['filters'].items() if v and k in labels) or 'Tüm kayıtlar'
    return title,filters,note,tables


def export_report(kind,context,format):
    title,filters,note,tables=report_tables(kind,context)
    output=BytesIO()
    if format=='xlsx':
        from openpyxl import Workbook
        from openpyxl.styles import Font,PatternFill,Alignment,Border,Side
        from openpyxl.utils import get_column_letter
        wb=Workbook();wb.remove(wb.active)
        for name,headers,rows,widths in tables:
            ws=wb.create_sheet(name)
            for row in [[title],[filters],[note],headers]+(rows or [['Kayıt bulunamadı.']]):
                ws.append([v.replace(tzinfo=None) if isinstance(v,datetime) else v for v in row])
            for row in ws:
                for cell in row:
                    if isinstance(cell.value,str):cell.data_type='s' # User text must never become an Excel formula.
                    cell.font=Font(name='Calibri',size=11,bold=cell.row in (1,4))
                    cell.alignment=Alignment(vertical='top',wrap_text=True)
                    if cell.row>=4:cell.border=Border(bottom=Side(style='thin',color='BCCADA'))
                    if cell.row==4:cell.fill=PatternFill('solid',fgColor='E8EEF5')
                    if isinstance(cell.value,datetime):cell.number_format='dd.mm.yyyy hh:mm'
                    elif isinstance(cell.value,date):cell.number_format='dd.mm.yyyy'
                    elif isinstance(cell.value,(float,Decimal)):cell.number_format='#,##0.00'
            for i,width in enumerate(widths,1):ws.column_dimensions[get_column_letter(i)].width=width
            for row in (1,2,3):ws.merge_cells(start_row=row,start_column=1,end_row=row,end_column=len(headers))
            ws.row_dimensions[3].height=32;ws.freeze_panes='A5';ws.auto_filter.ref=f'A4:{get_column_letter(len(headers))}{ws.max_row}'
            ws.print_title_rows='1:4';ws.page_setup.orientation='landscape';ws.page_setup.paperSize=ws.PAPERSIZE_A4;ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0;ws.sheet_properties.pageSetUpPr.fitToPage=True
        wb.save(output)
    elif format=='pdf':
        from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,Table,TableStyle,PageBreak
        from reportlab.lib.styles import ParagraphStyle
        from reportlab.lib.pagesizes import A4,landscape
        from reportlab.lib import colors
        from pdf_fonts import register_pdf_fonts
        regular,bold=register_pdf_fonts()
        body=ParagraphStyle('cell',fontName=regular,fontSize=7,leading=9,wordWrap='CJK')
        heading=ParagraphStyle('heading',fontName=bold,fontSize=14,leading=18)
        doc=SimpleDocTemplate(output,pagesize=landscape(A4),leftMargin=25,rightMargin=25,topMargin=25,bottomMargin=28,title=title)
        def display(v):
            if v is None:return '—'
            if isinstance(v,datetime):return v.strftime('%d.%m.%Y %H:%M')
            if isinstance(v,date):return v.strftime('%d.%m.%Y')
            if isinstance(v,(Decimal,float)):return f'{v:,.2f}'.replace(',','_').replace('.',',').replace('_','.')
            return str(v)
        para=lambda v:Paragraph(escape(display(v)),body)
        story=[]
        for index,(name,headers,rows,widths) in enumerate(tables):
            if index:story.append(PageBreak())
            story += [Paragraph(escape(title+' — '+name),heading),Spacer(1,8),para(filters),Spacer(1,5),para(note),Spacer(1,10)]
            table=Table([[para(v) for v in headers]]+[[para(v) for v in row] for row in rows],colWidths=[doc.width*w/sum(widths) for w in widths],repeatRows=1,hAlign='LEFT')
            table.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#e8eef5')),('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.white,colors.HexColor('#f1f4f8')]),('GRID',(0,0),(-1,-1),.4,colors.HexColor('#aebdce')),('VALIGN',(0,0),(-1,-1),'TOP'),('LEFTPADDING',(0,0),(-1,-1),5),('RIGHTPADDING',(0,0),(-1,-1),5)]))
            story.append(table)
            if not rows:story.append(para('Kayıt bulunamadı.'))
        def footer(canvas,doc):
            canvas.setFont(regular,8);canvas.drawRightString(landscape(A4)[0]-25,13,f'Sayfa {doc.page}')
        doc.build(story,onFirstPage=footer,onLaterPages=footer)
    else:raise ValueError('Unsupported report format')
    output.seek(0);return output
