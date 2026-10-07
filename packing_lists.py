"""Quotation-linked packing lists, with no stock or accounting side effects."""
import io,json
from collections import defaultdict
from datetime import date
from decimal import Decimal,InvalidOperation,ROUND_HALF_UP
from flask import request,render_template,redirect,url_for,send_file,abort,flash
from openpyxl import Workbook
from openpyxl.styles import Font,Alignment,Border,Side,PatternFill

META=('sender','sender_address','sender_phone','recipient','address','phone','reference','date','container','seal','notes')
HEADERS=['Marks','Model No','Item Name / HS Code','Country of Origin','Description','QTY (PCS)','CTN','G.W. (KG)','N.W. (KG)','CBM']
def number(value,label,integer=False):
    try:
        n=Decimal(str(value))
        if not n.is_finite() or n<0 or n>999999999 or (integer and n!=int(n)):raise ValueError()
        return n
    except (ValueError,TypeError,InvalidOperation):raise ValueError(label+': geçerli, sıfır veya pozitif sayı girin.')
def fixed(n,places=2):return str(n.quantize(Decimal(10)**-places,rounding=ROUND_HALF_UP))
def defaults(quote,source):
    return dict(sender='ABİKA MOBİLYA',sender_address='',sender_phone=source.get('author',{}).get('phone',''),
                recipient=source['customer_name'],address=source.get('address',''),phone=source.get('phone',''),reference=quote.number,
                date=date.today().isoformat(),container='',seal='',notes='',rows=[dict(product_id=r['product_id'],model=r['code'],name=r['name'],hs='',origin='',description=r.get('description',''),marks='',quantity=r['quantity'],cartons='',gross='',net='',cbm='',length='',width='',height='') for r in source['lines']])
def validate(data,source):
    if not isinstance(data,dict):raise ValueError('Packing list verisi geçersiz.')
    clean={k:str(data.get(k,''))[:1000].strip() for k in META}
    if not clean['sender'] or not clean['recipient']:raise ValueError('Gönderici ve alıcı adı zorunludur.')
    try:date.fromisoformat(clean['date'])
    except ValueError:raise ValueError('Belge tarihini kontrol edin.')
    rows=data.get('rows',[])
    if not isinstance(rows,list) or not 1<=len(rows)<=100:raise ValueError('1 ile 100 arasında paket satırı ekleyin.')
    available=defaultdict(int)
    for row in source['lines']:available[str(row['product_id'])]+=int(row['quantity'])
    used=defaultdict(int);clean['rows']=[]
    total={k:Decimal(0) for k in ('quantity','cartons','gross','net','cbm')}
    for row in rows:
        if not isinstance(row,dict) or str(row.get('product_id')) not in available:raise ValueError('Her satırda tekliften bir ürün seçin.')
        out={k:str(row.get(k,''))[:500].strip() for k in ('model','name','hs','origin','description','marks')}
        if not out['name'] or not out['origin']:raise ValueError('Ürün adı ve menşei zorunludur.')
        hs=out['hs']
        if hs and (not hs.isascii() or not hs.isdigit() or len(hs) not in (6,8,10,12)):raise ValueError('HS kodu 6, 8, 10 veya 12 rakam olmalıdır.')
        out['product_id']=int(row['product_id'])
        qty=number(row.get('quantity'),'Ürün adedi',True);ctn=number(row.get('cartons'),'Koli sayısı',True)
        if qty<1 or ctn<1:raise ValueError('Ürün adedi ve koli sayısı sıfırdan büyük olmalıdır.')
        used[str(row['product_id'])]+=int(qty)
        if used[str(row['product_id'])]>available[str(row['product_id'])]:raise ValueError(out['name']+': sevk adedi teklif adedini aşıyor.')
        gross=number(row.get('gross'),'Brüt ağırlık');net=number(row.get('net'),'Net ağırlık')
        if gross<=0 or net<=0 or gross<net:raise ValueError('Ağırlıklar sıfırdan büyük olmalı; brüt ağırlık net ağırlıktan küçük olamaz.')
        dimensions=[str(row.get(k,'')).strip() for k in ('length','width','height')]
        if any(dimensions):
            if not all(dimensions):raise ValueError('Otomatik hacim için üç paket ölçüsünü de girin.')
            dims=[number(v,'Paket ölçüsü') for v in dimensions]
            if any(v<=0 for v in dims):raise ValueError('Paket ölçüleri sıfırdan büyük olmalıdır.')
            cbm=dims[0]*dims[1]*dims[2]*ctn/Decimal(1000000)
        else:cbm=number(row.get('cbm'),'Toplam hacim')
        if cbm<=0:raise ValueError('Toplam hacim sıfırdan büyük olmalıdır.')
        out.update(quantity=int(qty),cartons=int(ctn),gross=fixed(gross),net=fixed(net),cbm=fixed(cbm,3),length=dimensions[0],width=dimensions[1],height=dimensions[2])
        if Decimal(out['net'])<=0 or Decimal(out['cbm'])<=0:raise ValueError('Ağırlık veya hacim belge hassasiyetinden küçük.')
        clean['rows'].append(out)
        for k in total:total[k]+=Decimal(str(out[k]))
    clean['totals']={k:int(v) if k in ('quantity','cartons') else fixed(v,3 if k=='cbm' else 2) for k,v in total.items()}
    return clean

def cells(row):
    return [row['marks'],row['model'],row['name']+('\n'+row['hs'] if row['hs'] else ''),row['origin'],row['description'],row['quantity'],row['cartons'],row['gross'],row['net'],row['cbm']]

def pdf(data,number):
    from reportlab.lib import colors
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.enums import TA_CENTER
    from reportlab.platypus import SimpleDocTemplate,Paragraph,Table,TableStyle,Spacer,KeepTogether
    from pdf_fonts import register_pdf_fonts
    from xml.sax.saxutils import escape
    font,bold=register_pdf_fonts();buf=io.BytesIO();navy=colors.HexColor('#19344D')
    style=ParagraphStyle('Packing',fontName=font,fontSize=8,leading=11)
    head=ParagraphStyle('PackingHead',parent=style,fontName=bold,alignment=TA_CENTER)
    def p(v,sty=style):return Paragraph(escape(str(v or '')).replace('\n','<br/>'),sty)
    doc=SimpleDocTemplate(buf,pagesize=(841.9,595.3),leftMargin=24,rightMargin=24,topMargin=24,bottomMargin=32)
    story=[p(data['sender'],ParagraphStyle('Company',parent=head,fontSize=17,leading=21,textColor=navy)),p(data['sender_address'],head),p('Tel: '+data['sender_phone'],head),Spacer(1,8),p('PACKING LIST',ParagraphStyle('Title',parent=head,fontSize=19,leading=24)),Spacer(1,10)]
    meta=Table([[p('To: '+data['recipient']),p('Invoice / Reference No: '+data['reference'])],[p('Address: '+data['address']),p('Date: '+data['date'])],[p('TEL: '+data['phone']),p('Container No: '+(data['container'] or '—'))],[p(''),p('Seal No: '+(data['seal'] or '—'))]],colWidths=[480,313.9])
    meta.setStyle(TableStyle([('GRID',(0,0),(-1,-1),.4,colors.grey),('VALIGN',(0,0),(-1,-1),'TOP')]));story.extend([meta,Spacer(1,12)])
    table_rows=[[p(h,head) for h in HEADERS]]
    for row in data['rows']:table_rows.append([p(v) for v in cells(row)])
    t=data['totals'];table_rows.append([p('TOTAL',head),'','','','',p(t['quantity'],head),p(t['cartons'],head),p(t['gross'],head),p(t['net'],head),p(t['cbm'],head)])
    table=Table(table_rows,colWidths=[48,75,110,62,173.9,52,44,75,75,79],repeatRows=1)
    table.setStyle(TableStyle([('GRID',(0,0),(-1,-1),.7,colors.black),('BACKGROUND',(0,0),(-1,0),colors.HexColor('#EDF4FA')),('VALIGN',(0,0),(-1,-1),'MIDDLE'),('TOPPADDING',(0,0),(-1,-1),9),('BOTTOMPADDING',(0,0),(-1,-1),9)]));story.append(table)
    story.extend([Spacer(1,12),p('Weights and CBM are totals per row. Dimensions used for CBM are external package dimensions in cm.')])
    if data['notes']:story.extend([Spacer(1,8),p(data['notes'])])
    story.extend([Spacer(1,20),p('Signature: __________________________',head)])
    def footer(c,d):
        if d.page==1:
            c.saveState();c.setFillColor(navy);c.rect(24,546,70,25,fill=1,stroke=0);c.setFillColor(colors.white);c.setFont(bold,16);c.drawCentredString(59,553,'ABİKA');c.restoreState()
        c.setFont(font,7);c.drawString(24,17,'ABİKA · '+number);c.drawRightString(817.9,17,'Page '+str(d.page))
    doc.build(story,onFirstPage=footer,onLaterPages=footer);buf.seek(0);return buf

def excel(data,number):
    wb=Workbook();ws=wb.active;ws.title='Packing List';line=Side(style='thin',color='333333')
    def text(value):
        # Keep supplier/customer-controlled values as text, never executable formulas.
        return "'"+value if isinstance(value,str) and value.startswith(('=','+','-','@')) else value
    for n,value in enumerate([data['sender'],data['sender_address'],'Tel: '+data['sender_phone'],'PACKING LIST'],1):
        ws.merge_cells(start_row=n,start_column=1,end_row=n,end_column=10);ws.cell(n,1,text(value));ws.cell(n,1).font=Font(size=18 if n in (1,4) else 10,bold=n in (1,4));ws.cell(n,1).alignment=Alignment(horizontal='center')
    for n,(left,right) in enumerate([('To: '+data['recipient'],'Invoice / Reference No: '+data['reference']),('Address: '+data['address'],'Date: '+data['date']),('TEL: '+data['phone'],'Container No: '+data['container']),('','Seal No: '+data['seal'])],5):
        ws.merge_cells(start_row=n,start_column=1,end_row=n,end_column=5);ws.merge_cells(start_row=n,start_column=6,end_row=n,end_column=10);ws.cell(n,1,text(left));ws.cell(n,6,text(right))
    for j,h in enumerate(HEADERS,1):ws.cell(10,j,h)
    for i,row in enumerate(data['rows'],11):
        for j,v in enumerate(cells(row),1):ws.cell(i,j,float(Decimal(str(v))) if j>=6 else text(v))
    end=11+len(data['rows']);ws.cell(end,1,'TOTAL')
    for j,k in enumerate(('quantity','cartons','gross','net','cbm'),6):ws.cell(end,j,float(Decimal(str(data['totals'][k]))))
    for row in ws.iter_rows(min_row=10,max_row=end,max_col=10):
        for c in row:c.border=Border(left=line,right=line,top=line,bottom=line);c.alignment=Alignment(vertical='center',wrap_text=True);c.font=Font(size=10,bold=c.row in (10,end))
        ws.row_dimensions[row[0].row].height=42 if row[0].row==10 else 54
    for c in ws[10]:c.fill=PatternFill('solid',fgColor='EDF4FA')
    for letter,width in zip('ABCDEFGHIJ',[12,18,25,15,38,12,10,15,15,14]):ws.column_dimensions[letter].width=width
    for row in ws.iter_rows(min_row=11,max_row=end,min_col=8,max_col=10):
        for c in row:c.number_format='0.000' if c.column==10 else '0.00'
    ws.cell(end+2,1,text(data['notes']));ws.cell(end+4,5,'Signature:');ws.freeze_panes='F11';ws.print_title_rows='1:10';ws.sheet_properties.pageSetUpPr.fitToPage=True;ws.page_setup.orientation='landscape';ws.page_setup.paperSize=ws.PAPERSIZE_A4;ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0;ws.print_options.horizontalCentered=True;ws.print_area=f'A1:J{end+4}'
    buf=io.BytesIO();wb.save(buf);buf.seek(0);return buf

def register_packing(app,db,Quote):
    def load(uid):
        q=db.get_or_404(Quote,uid);return q,json.loads(q.payload)
    @app.route('/fiyat-teklifleri/<uid>/packing-list',methods=['GET','POST'])
    def packing_edit(uid):
        q,source=load(uid);data=source.get('packing_list') or defaults(q,source);error=None
        if request.method=='POST':
            try:
                if request.form.get('version')!=str(q.version):raise ValueError('Teklif veya packing list başka bir oturumda değişti. Sayfayı yenileyin.')
                data=json.loads(request.form.get('payload','{}'));clean=validate(data,source)
                source['packing_list']=clean
                changed=Quote.query.filter_by(id=q.id,version=q.version).update(dict(payload=json.dumps(source,ensure_ascii=False),version=q.version+1),synchronize_session=False)
                if changed!=1:raise ValueError('Kayıt değişti; sayfayı yenileyip kontrol edin.')
                db.session.commit();flash('Packing list kaydedildi. PDF ve Excel indirebilirsiniz.','success');return redirect(url_for('packing_edit',uid=q.id))
            except (ValueError,TypeError,KeyError) as e:
                db.session.rollback();error=str(e)
                if not isinstance(data,dict) or not isinstance(data.get('rows',[]),list):data=defaults(q,source)
        return render_template('packing_edit.html',quote=q,source=source,data=data,error=error,saved=bool(source.get('packing_list')),meta=META,products=[{k:r[k] for k in ('product_id','code','name','description','quantity')} for r in source['lines']])
    @app.get('/fiyat-teklifleri/<uid>/packing-list/<kind>')
    def packing_export(uid,kind):
        q,source=load(uid)
        if kind not in ('pdf','xlsx') or not source.get('packing_list'):abort(404)
        try:data=validate(source['packing_list'],source)
        except ValueError as e:abort(400,str(e))
        buf=pdf(data,q.number) if kind=='pdf' else excel(data,q.number)
        return send_file(buf,as_attachment=True,download_name=q.number+'-Packing-List.'+kind,mimetype='application/pdf' if kind=='pdf' else 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
