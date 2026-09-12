"""Isolated tests. Optional generated fixtures go only to a temporary QA directory."""
import os
import tempfile
import unittest
from datetime import date
from decimal import Decimal
from io import BytesIO
from pathlib import Path

_data = tempfile.TemporaryDirectory(prefix='businessos-statement-test-')
os.environ['BUSINESSOS_DATA_DIR'] = _data.name
os.environ['DATABASE_URL'] = 'sqlite:///:memory:'
from app import app, db, Customer, Order, OrderItem, Product, Invoice, InvoiceItem, AccountTransaction, build_account_statement
from account_exports import statement_period
from desktop import EXPORT_PATH
from openpyxl import load_workbook


class StatementExportTests(unittest.TestCase):
    def setUp(self):
        self.context = app.app_context()
        self.context.push()
        db.drop_all()
        db.create_all()
        self.customer = Customer(name='İNCİ MOBİLYA İNŞAAT SANAYİ VE TİCARET LİMİTED ŞİRKETİ', code='C01')
        other = Customer(name='Diğer müşteri')
        product = Product(name='Ürün')
        db.session.add_all([self.customer, other, product])
        db.session.flush()
        order = Order(customer=self.customer, order_no='SS-TEST', order_type='Satış', status='Teslim Edildi', notes='Sipariş özel notu')
        order.items.append(OrderItem(product_name='Stoksuz hizmet', quantity=1, unit_price=999, vat_rate=0, variant='Gri', detail_2='Özel montaj', detail_3='Kollu', note='Kalem notu'))
        sale = Invoice(customer=self.customer, invoice_no='F-SALE', invoice_type='Satış', invoice_date=date(2026,9,2), order=order, notes='Fatura notu', due_date=date(2026,9,30))
        sale.items.append(InvoiceItem(product=product, product_name='=Özel Ürün ŞĞİıÇ & montaj', quantity=2, unit_price=100, discount_rate=10, vat_rate=10))
        purchase = Invoice(customer=self.customer, invoice_no='F-BUY', invoice_type='Satın Alma', invoice_date=date(2026,9,3))
        purchase.items.append(InvoiceItem(product=product, product_name='Yedek parça', quantity=1, unit_price=40, vat_rate=0))
        db.session.add_all([sale,purchase])
        for customer, day, debit, credit, ref in [(self.customer,date(2026,8,1),300,0,'OPEN'),(self.customer,date(2026,9,4),0,50,'PAY'),(self.customer,date(2026,10,1),777,0,'FUTURE'),(other,date(2026,9,2),888,0,'OTHER')]:
            db.session.add(AccountTransaction(customer=customer,transaction_date=day,transaction_type='Tahsilat',description='Ödeme açıklaması',debit=debit,credit=credit,reference_no=ref,payment_method='Çek',check_no='C123',check_bank='Test Bankası',check_due_date=date(2026,10,2),check_status='Bekliyor'))
        db.session.commit()
        self.client = app.test_client()
        self.path = f'/musteriler/{self.customer.id}/cari-hesap'
        self.query = '?start_date=2026-09-01&end_date=2026-09-30'

    def tearDown(self):
        db.session.remove()
        self.context.pop()

    def test_period_and_screen_links(self):
        period = statement_period(build_account_statement(self.customer), date(2026,9,1),date(2026,9,30))
        self.assertEqual(period['opening_balance'],300)
        self.assertEqual(period['period_debit'],198)
        self.assertEqual(period['period_credit'],90)
        self.assertEqual(period['closing_balance'],408)
        self.assertEqual(period['entries'][-1]['balance'],408)
        html = self.client.get(self.path+self.query).get_data(as_text=True)
        for detail in ['ozet','ayrintili']:
            for fmt in ['excel','pdf']:
                path = self.path+f'/ekstre/{detail}/{fmt}'
                self.assertIn(path,html)
                self.assertTrue(EXPORT_PATH.fullmatch(path))
        self.assertIn('end_date=2026-09-30',html)
        self.assertEqual(self.client.get(self.path+'/ekstre/invalid/pdf').status_code,404)

    def test_excel_summary_and_details(self):
        for detail in ['ozet','ayrintili']:
            response = self.client.get(self.path+f'/ekstre/{detail}/excel'+self.query)
            self.assertEqual(response.status_code,200)
            self.assertIn('attachment', response.headers['Content-Disposition'])
            book = load_workbook(BytesIO(response.data))
            ws = book['Ekstre']
            self.assertEqual(ws['F5'].value,300)
            self.assertEqual(ws['F9'].value,408)
            self.assertEqual(ws['D9'].value,198)
            self.assertEqual(ws['E9'].value,90)
            self.assertNotIn('FUTURE',str(list(ws.values)))
            self.assertNotIn('OTHER',str(list(ws.values)))
            if detail == 'ayrintili':
                self.assertEqual(len(book.sheetnames),4)
                self.assertEqual(book['Fatura Kalemleri']['D5'].data_type,'s')
                self.assertEqual(book['Fatura Kalemleri']['P5'].value,198)
                self.assertEqual(book['Bağlı Sipariş Kalemleri']['F5'].value,'Özel montaj')
                self.assertIn('Test Bankası',book['Ödeme ve Tahsilat']['D5'].value)
            else:
                self.assertEqual(len(book.sheetnames),1)
            if os.environ.get('STATEMENT_QA_DIR'):
                Path(os.environ['STATEMENT_QA_DIR'],f'{detail}.xlsx').write_bytes(response.data)
        empty = load_workbook(BytesIO(self.client.get(self.path+'/ekstre/ozet/excel?start_date=2026-09-10&end_date=2026-09-15').data))
        self.assertEqual(empty['Ekstre']['F5'].value,408)
        self.assertEqual(empty['Ekstre']['F6'].value,408)

    def test_pdf_and_multipage(self):
        # Enough rows to exercise pagination, headers and Turkish text wrapping.
        for index in range(30):
            db.session.add(AccountTransaction(customer=self.customer, transaction_date=date(2026,9,5), transaction_type='Ödeme', debit=0, credit=0, reference_no=f'REF-{index}', description='Türkçe açıklama: Şişli, Üsküdar, ölçü ve özel montaj. '*5, payment_method='Kredi Kartı', card_installments=3, card_owner_type='Müşteri Kartı', card_customer_name='İnci'))
        db.session.commit()
        for detail in ['ozet','ayrintili']:
            response = self.client.get(self.path+f'/ekstre/{detail}/pdf'+self.query)
            self.assertEqual(response.status_code,200)
            self.assertTrue(response.data.startswith(b'%PDF-'))
            if os.environ.get('STATEMENT_QA_DIR'):
                Path(os.environ['STATEMENT_QA_DIR'],f'{detail}.pdf').write_bytes(response.data)


if __name__ == '__main__':
    unittest.main()
