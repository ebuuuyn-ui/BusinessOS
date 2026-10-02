import unittest
from io import BytesIO
from datetime import date
from decimal import Decimal
from openpyxl import load_workbook
from test_supplier_sheets_routes import SupplierRouteTests
from test_web_auth import BASE
import app as m

class InvoiceExportTests(unittest.TestCase):
    setUp = SupplierRouteTests.setUp
    tearDown = SupplierRouteTests.tearDown

    def test_export_all_types_dates_amounts_and_safe_text(self):
        customer=m.db.session.get(m.Customer,1339)
        customer.name='=DANGEROUS()'; customer.tax_number='0123456789'
        for kind in ('Satış','Satın Alma','Satış İadesi'):
            invoice=m.Invoice(invoice_no=kind,invoice_type=kind,customer=customer,
                order_id=self.order_id if kind=='Satış' else None,
                invoice_date=date(2026,9,30),due_date=date(2026,10,30))
            invoice.items=[m.InvoiceItem(product_name='Test',quantity=2,unit_price=Decimal('100'),vat_rate=10)]
            m.db.session.add(invoice)
        m.db.session.commit()
        result=self.client.get('/faturalar/excel?type=Satış',base_url=BASE)
        self.assertEqual(result.status_code,200)
        self.assertIn('attachment',result.headers['Content-Disposition'])
        sheet=load_workbook(BytesIO(result.data)).active
        self.assertEqual(sheet.max_row,4)
        self.assertEqual({sheet.cell(i,4).value for i in range(2,5)},{'Satış','Satın Alma','Satış İadesi'})
        self.assertEqual(sheet['H2'].value,'0123456789')
        self.assertEqual(sheet['F2'].data_type,'s')
        self.assertEqual(sheet['K2'].value,200)
        self.assertEqual(sheet['L2'].value,20)
        self.assertEqual(sheet['M2'].value,220)
        self.assertEqual(sheet['A2'].value.date(),date(2026,9,30))
        self.assertEqual(sheet.freeze_panes,'A2')
        self.assertEqual(self.app.test_client().get('/faturalar/excel',base_url=BASE).status_code,302)

    def test_empty_register_and_button(self):
        result=self.client.get('/faturalar/excel',base_url=BASE)
        self.assertEqual(load_workbook(BytesIO(result.data)).active.max_row,1)
        self.assertIn('/faturalar/excel',self.client.get('/faturalar',base_url=BASE).text)
