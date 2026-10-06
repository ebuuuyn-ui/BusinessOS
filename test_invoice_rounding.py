"""Supplier penny rounding must not change official totals or stock quantities."""
import copy
import re
import unittest
from decimal import Decimal as D
from unittest.mock import patch
from sqlalchemy import create_engine, text, inspect
import test_uyumsoft_inbox as f
import app as m
from test_web_auth import BASE
from uyumsoft_client import A, B, UyumError
from uyumsoft_inbox import parse_inbox
from invoice_amounts import prepare_schema, schema_ready
from invoice_export import build_invoice_register
from openpyxl import load_workbook


class RoundingTests(unittest.TestCase):
    setUp=f.InboxTests.setUp
    tearDown=f.InboxTests.tearDown
    fixture=f.InboxTests.fixture
    make_invoice=f.InboxTests.make_invoice
    connect=f.InboxTests.connect
    post=f.InboxTests.post
    mock=f.InboxTests.mock

    def example(self):
        self.fixture()
        line=self.root.find('{'+A+'}InvoiceLine')
        second=copy.deepcopy(line);self.root.append(second)
        for i, row in enumerate((line,second)):
            row.find('{'+B+'}ID').text=str(i+1)
            row.find('{'+B+'}InvoicedQuantity').text='20'
            row.find('{'+B+'}InvoicedQuantity').set('unitCode','NIU')
            row.find('{'+B+'}LineExtensionAmount').text='27272.73'
            row.find('{'+A+'}Item/{'+B+'}Name').text='SİSA '+('SOMON' if not i else 'GRİ')
            row.find('{'+A+'}Price/{'+B+'}PriceAmount').text='1363.63636'
            for x in row.findall('{'+A+'}AllowanceCharge'):row.remove(x)
            for x in row.findall('.//{'+B+'}TaxAmount'):x.text='2727.28' if not i else '2727.27'
        self.root.find('{'+A+'}TaxTotal/{'+B+'}TaxAmount').text='5454.55'
        monetary=self.root.find('{'+A+'}LegalMonetaryTotal')
        for name,val in [('LineExtensionAmount','54545.45'),('TaxExclusiveAmount','54545.45'),('TaxInclusiveAmount','60000.00'),('PayableAmount','60000.00')]:
            monetary.find('{'+B+'}'+name).text=val
        self.form.update(kind_1='stock',product_1=self.form['product_0'])
        order=m.db.session.get(m.Order,self.order_id)
        order.items=[m.OrderItem(product=self.product,product_name='SİSA',quantity=40,unit='Adet')]
        m.db.session.commit()
        self.form['order_refs']=order.order_no
        for i in range(2):
            self.form[f'allocation_item_{i}[]']=str(order.items[0].id)
            self.form[f'allocation_qty_{i}[]']='20'

    def test_exact_official_totals_survive_save_reload_export_and_ledger(self):
        self.example()
        with patch('uyumsoft_inbox.Client') as client:
            self.mock(client)
            page=self.client.get(self.path,base_url=BASE)
            self.assertIn('kuruş yuvarlama',page.text)
            self.assertIn('1.363,63636',page.text)
            self.form['token']=re.search(r'name="token" value="([^"]+)"',page.text)[1]
            # Preview must bind the requested order for the final submission.
            page=self.client.get(self.path+'?order_refs=SA-TEST',base_url=BASE)
            self.form['token']=re.search(r'name="token" value="([^"]+)"',page.text)[1]
            self.assertEqual(m.Invoice.query.count(),0)
            result=self.post(self.path,self.form)
            self.assertEqual(result.status_code,302,result.text)
            self.assertEqual(self.post(self.path,self.form).status_code,302)
        m.db.session.remove()
        invoice=m.Invoice.query.one()
        self.assertEqual(invoice.net_amount,D('54545.45'))
        self.assertEqual(invoice.vat_amount,D('5454.55'))
        self.assertEqual(invoice.total_amount,D('60000.00'))
        self.assertEqual(invoice.items[0].unit_price,D('1363.63636'))
        self.assertEqual(invoice.items[0].vat_amount,D('2727.28'))
        self.assertEqual(invoice.rounding_difference,D('-.01'))
        self.assertEqual(invoice.customer.balance,D('-60000'))
        self.assertEqual(m.load_invoice_ledger(customer_id=1339)[1][0].total_amount,D('60000'))
        self.assertEqual(sum(x.quantity for x in m.StockMovement.query.all()),40)
        self.assertEqual(m.InvoiceOrderAllocation.query.count(),2)
        sheet=load_workbook(build_invoice_register([invoice])).active
        self.assertEqual(sheet.cell(2,13).value,60000)
        detail=self.client.get(f'/faturalar/{invoice.id}',base_url=BASE)
        self.assertIn('Fatura Geneli Yuvarlama Farkı',detail.text)
        self.assertIn('1.363,63636',detail.text)
        self.assertEqual(self.post(f'/faturalar/{invoice.id}/duzenle',{}).status_code,302)
        self.assertEqual(m.Invoice.query.one().total_amount,D('60000'))

    def test_larger_differences_and_wrong_header_totals_rejected(self):
        self.example()
        targets=[('{'+A+'}InvoiceLine/{'+A+'}TaxTotal/{'+A+'}TaxSubtotal/{'+B+'}TaxAmount','2727.29'),
                 ('{'+A+'}LegalMonetaryTotal/{'+B+'}TaxExclusiveAmount','54545.00'),
                 ('{'+A+'}LegalMonetaryTotal/{'+B+'}PayableAmount','60000.01')]
        for path,value in targets:
            root=copy.deepcopy(self.root);root.find(path).text=value
            with self.assertRaises(UyumError):parse_inbox(root,self.uid,f.fixture.DEFAULT_SELLER['tax_number'])
        self.assertEqual(m.Invoice.query.count(),0)
        self.assertEqual(m.StockMovement.query.count(),0)

    def test_manual_invoices_keep_existing_calculations(self):
        invoice=self.make_invoice()
        self.assertEqual(invoice.total_amount,D('198'))
        self.assertIsNone(invoice.source_net_amount)


class SchemaTests(unittest.TestCase):
    def test_additive_and_idempotent_sqlite_upgrade(self):
        engine=create_engine('sqlite:///:memory:')
        with engine.begin() as c:
            c.execute(text('CREATE TABLE invoice (id INTEGER PRIMARY KEY, notes TEXT)'))
            c.execute(text('CREATE TABLE invoice_item (id INTEGER PRIMARY KEY, unit_price NUMERIC(12,2))'))
            c.execute(text("INSERT INTO invoice VALUES (1, 'preserve')"))
            c.execute(text('INSERT INTO invoice_item VALUES (1, 1363.64)'))
        prepare_schema(engine);prepare_schema(engine)
        with engine.connect() as c:
            self.assertTrue(schema_ready(c))
            self.assertEqual(c.execute(text('SELECT notes FROM invoice')).scalar(),'preserve')
            self.assertEqual(str(c.execute(text('SELECT unit_price FROM invoice_item')).scalar()),'1363.64')
            self.assertIsNone(c.execute(text('SELECT source_net_amount FROM invoice')).scalar())
        engine.dispose()

if __name__=='__main__':unittest.main()
