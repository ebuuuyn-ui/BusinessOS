"""Excel/PDF customer summaries from the same context as the tracking screen."""
from maturity_labels import tracking_text
from datetime import date
from decimal import Decimal
from io import BytesIO
from xml.sax.saxutils import escape
from collection_exports import filter_label

HEADERS = ['Cari', 'Açık Kayıt', 'Net Bakiye', 'Gecikmiş Tutar', 'Ortalama Vade', 'Vadesiz Fatura Tutarı', 'En Eski Açık Vade', 'Durum']
NOTE = ('Net bakiye faturalar ve tüm cari hareketleriyle aynıdır. Mahsuplar en eski fatura/hareketten düşülür. '
        'Ortalama vade yalnız vadesi bilinen açık tutarlardan hesaplanır. Vadesiz faturalar ortalamaya ve gecikmeye dahil edilmez. Fatura dışı bakiye net tutara dahildir; vade uygulanmaz. '
        'Arama cariyi seçer; tüm hareketleri hesapta kalır. Tutarlar TL.')


def rows(context):
    return [[g['customer'].name,g['open_count'],g['remaining'],g['overdue_amount'],g['average_due_date'],
             g['undated_amount'],g['oldest_due_date'],g['state_label']] for g in context['customers']]


def total_row(context):
    groups=context['customers']
    return ['Liste Toplamı',sum(g['open_count'] for g in groups),sum((g['remaining'] for g in groups),Decimal('0')),
            sum((g['overdue_amount'] for g in groups),Decimal('0')),None,sum((g['undated_amount'] for g in groups),Decimal('0')),None,None]


def export_excel(context):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    book=Workbook()
    ws=book.active
    ws.title=tracking_text('Cari Tahsilat Özeti', context)
    ws.append([tracking_text('Cari Bazında Tahsilat Takibi', context)])
    ws.append([f"Rapor tarihi: {context['today']:%d.%m.%Y} | {filter_label(context)}"])
    ws.append([tracking_text(NOTE, context)])
    ws.append([f"Filtrelenen liste: {len(context['customers'])} cari"])
    for n in range(1,5): ws.merge_cells(start_row=n,start_column=1,end_row=n,end_column=8)
    ws.append([tracking_text(h, context) for h in HEADERS])
    for row in rows(context): ws.append([float(v) if isinstance(v,Decimal) else v for v in row])
    last_data=ws.max_row
    ws.append([float(v) if isinstance(v,Decimal) else v for v in total_row(context)])
    widths=[48,14,22,22,20,18,22,25]
    for i,width in enumerate(widths,1): ws.column_dimensions[get_column_letter(i)].width=width
    for row in ws:
        for cell in row:
            if cell.data_type=='f': cell.data_type='s'
            cell.font=Font(name='Arial',size=10,color='172033')
            cell.alignment=Alignment(vertical='top',wrap_text=True)
            if isinstance(cell.value,date): cell.number_format='dd.mm.yyyy'
            if cell.row>=6 and cell.column in (3,4,6):
                cell.number_format='#,##0.00'
                cell.alignment=Alignment(horizontal='right',vertical='top')
            if cell.row in (5,ws.max_row):
                cell.fill=PatternFill('solid',fgColor='24456A' if cell.row==5 else 'E9EFF8')
                cell.font=Font(name='Arial',size=10,bold=True,color='FFFFFF' if cell.row==5 else '172033')
        ws.row_dimensions[row[0].row].height=max(32,14*((len(str(row[0].value or ''))+42)//43)) if row[0].row>=6 else 32
    ws.row_dimensions[3].height=54
    ws['A1'].font=Font(name='Arial',size=16,bold=True,color='172033')
    ws.freeze_panes='B6'
    ws.auto_filter.ref=f'A5:H{last_data}'
    ws.print_title_rows='5:5'
    ws.sheet_view.showGridLines=False
    ws.page_setup.orientation='landscape'
    ws.page_setup.paperSize=ws.PAPERSIZE_A4
    ws.sheet_properties.pageSetUpPr.fitToPage=True
    ws.page_setup.fitToWidth=1
    ws.page_setup.fitToHeight=0
    out=BytesIO();book.save(out);out.seek(0)
    return out


def export_pdf(context):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    from pdf_fonts import register_pdf_fonts
    register_pdf_fonts('CollectionSummary', 'CollectionSummaryBold')
    body=ParagraphStyle('csbody',fontName='CollectionSummary',fontSize=8,leading=11,wordWrap='CJK')
    title=ParagraphStyle('cstitle',parent=body,fontName='CollectionSummaryBold',fontSize=16,leading=21,spaceAfter=9)
    header=ParagraphStyle('cshead',parent=body,fontName='CollectionSummaryBold',textColor=colors.white)
    right=ParagraphStyle('csright',parent=body,alignment=2)
    p=lambda v,style=body:Paragraph(escape(str(v)).replace('\n','<br/>'),style)
    money=lambda v:f'{v:,.2f}'.replace(',','X').replace('.',',').replace('X','.')
    width=landscape(A4)[0]-48
    out=BytesIO()
    doc=SimpleDocTemplate(out,pagesize=landscape(A4),leftMargin=24,rightMargin=24,topMargin=24,bottomMargin=30,title=tracking_text('Cari Bazında Tahsilat Takibi', context))
    story=[p(tracking_text('Cari Bazında Tahsilat Takibi', context),title),p(f"Rapor tarihi: {context['today']:%d.%m.%Y} | {filter_label(context)}"),Spacer(1,6),p(tracking_text(NOTE, context)),Spacer(1,9),p(f"Filtrelenen liste: {len(context['customers'])} cari"),Spacer(1,6)]
    data=[[p('Açık\nKayıt' if h=='Açık Kayıt' else h,header) for h in [tracking_text(h, context) for h in HEADERS]]]
    for row in rows(context)+[total_row(context)]:
        data.append([p('—' if v is None else money(v) if idx in (2,3,5) else v.strftime('%d.%m.%Y') if isinstance(v,date) else v,right if idx in (1,2,3,5) else body) for idx,v in enumerate(row)])
    table=Table(data,colWidths=[width*x/100 for x in (27,7,12,12,11,8,12,11)],repeatRows=1,splitByRow=1,splitInRow=1)
    table.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#24456A')),('BACKGROUND',(0,-1),(-1,-1),colors.HexColor('#E9EFF8')),('VALIGN',(0,0),(-1,-1),'TOP'),('LINEBELOW',(0,0),(-1,-1),.3,colors.HexColor('#D6DFEA')),('ROWBACKGROUNDS',(0,1),(-1,-2),[colors.white,colors.HexColor('#F4F7FB')]),('LEFTPADDING',(0,0),(-1,-1),5),('RIGHTPADDING',(0,0),(-1,-1),5),('TOPPADDING',(0,0),(-1,-1),6),('BOTTOMPADDING',(0,0),(-1,-1),6)]))
    story.append(table)
    if not context['customers']: story.append(p('Bu filtreye uygun cari bulunmuyor.'))
    def footer(canvas,doc):
        canvas.saveState();canvas.setFont('CollectionSummary',8);canvas.drawString(24,14,tracking_text('Business OS | Cari Tahsilat Özeti | Tutarlar TL', context));canvas.drawRightString(width+24,14,f'Sayfa {doc.page}');canvas.restoreState()
    doc.build(story,onFirstPage=footer,onLaterPages=footer);out.seek(0)
    return out
