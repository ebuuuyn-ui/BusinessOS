"""Invoice maturity must reconcile with the same ledger as customer balances."""
import os
import tempfile
import unittest
from datetime import date, timedelta
from decimal import Decimal
from io import BytesIO
from openpyxl import load_workbook

_data=tempfile.TemporaryDirectory(prefix='businessos-invoice-maturity-')
os.environ['BUSINESSOS_DATA_DIR']=_data.name
os.environ['DATABASE_URL']='sqlite:///:memory:'
os.environ['BUSINESSOS_BACKUP_DIR']=_data.name
from app import app, db, Customer, Product, Invoice, InvoiceItem, AccountTransaction, Order, OrderItem, calculate_customer_balances
from invoice_maturity import build_maturity_groups


class InvoiceMaturityTests(unittest.TestCase):
    def setUp(self):
        self.ctx=app.app_context();self.ctx.push();db.drop_all();db.create_all()
        self.today=date.today()
        self.customer=Customer(name='Aynı Cari',code='C-1')
        self.other=Customer(name='Devir Cari',code='C-2')
        self.product=Product(name='Test Ürünü')
        db.session.add_all([self.customer,self.other,self.product]);db.session.flush()
        self.client=app.test_client()

    def tearDown(self):
        db.session.remove();self.ctx.pop()

    def invoice(self,number,kind,amount,days=None,age=0):
        i=Invoice(customer=self.customer,invoice_no=number,invoice_type=kind,
                  invoice_date=self.today-timedelta(days=age),due_date=self.today+timedelta(days=days) if days is not None else None)
        i.items.append(InvoiceItem(product=self.product,product_name='Ürün',quantity=1,unit_price=amount,vat_rate=0))
        db.session.add(i);db.session.flush();return i

    def tx(self,debit,credit,customer=None):
        db.session.add(AccountTransaction(customer=customer or self.customer,transaction_type='Devir',description='Test',debit=debit,credit=credit,transaction_date=self.today))
        db.session.flush()

    def groups(self,purchase=False):
        return build_maturity_groups(Customer.query.all(),Invoice.query.all(),AccountTransaction.query.all(),self.today,purchase)

    def test_purchase_balance_matches_accounts_with_cross_invoices_and_adjustments(self):
        self.invoice('AL-1','Satın Alma',1000,days=-5,age=40)
        self.invoice('AL-2','Satın Alma',500,days=None,age=20)
        self.invoice('SAT-1','Satış',200,days=10)
        self.tx(400,0);self.tx(0,50)
        self.tx(0,75,self.other)
        order=Order(customer=self.customer,order_no='UNBILLED',order_type='Satın Alma',status='Teslim Edildi')
        order.items.append(OrderItem(product_name='Faturasız',quantity=1,unit_price=99999,vat_rate=0));db.session.add(order);db.session.commit()
        g=self.groups(True);balances=calculate_customer_balances()
        self.assertEqual(sum(x['remaining'] for x in g),Decimal('1025'))
        for x in g:self.assertEqual(x['remaining'],-balances[x['customer'].id])
        party=next(x for x in g if x['customer'].id==self.customer.id)
        self.assertEqual(party['remaining'],950)
        self.assertEqual(party['overdue_amount'],400)
        self.assertEqual(party['undated_amount'],500)
        self.assertEqual(party['average_due_date'],self.today-timedelta(days=5))
        self.assertEqual(self.groups(False),[])

    def test_sales_balance_unknown_dates_and_no_order_influence(self):
        self.invoice('SAT-1','Satış',300,days=None,age=70)
        self.invoice('SAT-2','Satış',700,days=5,age=1)
        self.tx(0,100);db.session.commit()
        g=self.groups()[0]
        self.assertEqual(g['remaining'],calculate_customer_balances()[self.customer.id])
        self.assertEqual(g['remaining'],900);self.assertEqual(g['undated_amount'],200)
        self.assertEqual(g['due_soon_amount'],700);self.assertEqual(g['overdue_amount'],0)
        self.assertEqual(g['average_due_date'],self.today+timedelta(days=5))

    def test_overpayment_moves_to_opposite_side_without_invented_due_date(self):
        self.invoice('SAT','Satış',100,days=-50)
        self.tx(0,150);db.session.commit()
        self.assertEqual(self.groups(),[])
        g=self.groups(True)[0]
        self.assertEqual(g['remaining'],50);self.assertEqual(g['undated_amount'],0);self.assertEqual(g['other_amount'],50)
        self.assertEqual(g['overdue_amount'],0);self.assertIsNone(g['average_due_date'])

    def test_receipts_and_payments_have_no_maturity_warning(self):
        from invoice_maturity import filter_groups
        from collection_exports import rows
        for purchase, kind, debit, credit in [(True, 'Tahsilat', 0, 150), (False, 'Ödeme', 150, 0)]:
            with self.subTest(kind=kind):
                db.session.query(AccountTransaction).delete()
                db.session.add(AccountTransaction(customer=self.customer,transaction_type=kind,
                    description=kind,debit=debit,credit=credit,transaction_date=self.today))
                db.session.commit()
                g=self.groups(purchase)[0]
                self.assertEqual(g['remaining'],150)
                self.assertEqual(g['undated_amount'],0)
                self.assertEqual(g['other_amount'],150)
                self.assertEqual(g['average_label'],'—')
                self.assertEqual(g['entries'][0]['state'],'other')
                self.assertEqual(filter_groups([g],'undated','',str.lower),[])
                self.assertEqual(filter_groups([g],'other','',str.lower),[g])
                item=dict(g['entries'][0],customer=self.customer)
                self.assertEqual(rows(dict(items=[item]))[0][3],'—')
                url='/tahsilat-takibi?kind='+('purchase' if purchase else 'sales')
                html=self.client.get('/tahsilat-takibi/cari/'+str(self.customer.id)+'/ayrinti?kind='+('purchase' if purchase else 'sales')).get_data(as_text=True)
                detail=html.split('Fatura dışı cari hareketleri · Vade uygulanmaz')[1]
                self.assertNotIn('Vadesi belirtilmemiş',detail)
                self.assertIn('Fatura dışı bakiye',detail)
                response=self.client.get('/tahsilat-takibi/excel',query_string=dict(kind='purchase' if purchase else 'sales',view='details'))
                ws=load_workbook(BytesIO(response.data)).active
                self.assertEqual(ws['D10'].value,'—')
                self.assertEqual(ws['H10'].value,'Fatura dışı bakiye')

    def test_future_movements_and_zero_balance_match_account_screen(self):
        self.invoice('SAT','Satış',100,days=10)
        self.tx(0,100)
        AccountTransaction.query.one().transaction_date=self.today+timedelta(days=5)
        db.session.commit()
        self.assertEqual(self.groups()[0]['remaining'],calculate_customer_balances()[self.customer.id])
        self.assertEqual(self.groups()[0]['state'],'paid')

    def test_invoice_math_uses_vat_discount_and_multiple_invoices(self):
        i=self.invoice('SAT-1','Satış',100,days=0)
        i.items[0].quantity=2;i.items[0].discount_rate=10;i.items[0].vat_rate=20
        self.invoice('SAT-2','Satış',50,days=12);db.session.commit()
        self.assertEqual(self.groups()[0]['remaining'],Decimal('266.00'))

    def test_both_tabs_exports_filter_and_read_only(self):
        self.invoice('AL-1','Satın Alma',100,days=-2,age=40)
        self.invoice('AL-2','Satın Alma',200,days=None)
        self.tx(0,50,self.other);db.session.commit()
        def snapshot():
            raw=db.engine.raw_connection()
            try:return '\n'.join(raw.iterdump())
            finally:raw.close()
        before=snapshot()
        html=self.client.get('/tahsilat-takibi?kind=purchase&q=AL-2').get_data(as_text=True)
        self.assertIn('Vadesi Belirtilmemiş',html)
        self.assertNotIn('AL-1',html)
        detail=self.client.get('/tahsilat-takibi/cari/'+str(self.customer.id)+'/ayrinti?kind=purchase').get_data(as_text=True)
        self.assertIn('AL-1',detail);self.assertIn('AL-2',detail)
        self.assertNotIn('30 gün',html);self.assertNotIn('Teslim edilmiş',html)
        self.assertIn('/faturalar/',detail)
        self.assertIn('Vadesi belirtilmemiş',html)
        for view in ['summary','details']:
            r=self.client.get('/tahsilat-takibi/excel',query_string=dict(kind='purchase',q='AL-2',view=view))
            self.assertEqual(r.status_code,200)
            ws=load_workbook(BytesIO(r.data)).active
            self.assertIn('Ödeme',ws.title)
            if view=='summary':
                self.assertEqual(ws['C6'].value,300);self.assertEqual(ws['F6'].value,200)
            else:
                values=[c.value for row in ws for c in row]
                self.assertIn('AL-1',values);self.assertIn('AL-2',values);self.assertIn('Vadesi belirtilmemiş',values)
            r=self.client.get('/tahsilat-takibi/pdf',query_string=dict(kind='purchase',view=view))
            self.assertEqual(r.status_code,200);self.assertTrue(r.data.startswith(b'%PDF'))
        self.assertEqual(self.client.get('/tahsilat-takibi?kind=sales').status_code,200)
        for state in ['open','paid','overdue','due_soon','undated','other','all']:
            self.assertEqual(self.client.get('/tahsilat-takibi',query_string=dict(kind='purchase',state=state)).status_code,200)
        self.assertEqual(before,snapshot())

if __name__=='__main__':unittest.main()
