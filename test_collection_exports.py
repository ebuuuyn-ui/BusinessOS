"""Shared exports use invoice references, nullable due dates and safe text."""
import unittest
from datetime import date
from decimal import Decimal
from types import SimpleNamespace
from io import BytesIO
from openpyxl import load_workbook
from collection_exports import export_excel, export_pdf

class CollectionExportTests(unittest.TestCase):
    def context(self, count):
        item=dict(invoice=SimpleNamespace(id=1),customer=SimpleNamespace(name='=FORMULA <Firma>'),reference='F-001',document_date=date.today(),due_date=None,
                  amount=Decimal('100'),collected=Decimal('0'),remaining=Decimal('100'),state_label='Vadesi belirtilmemiş')
        return dict(is_purchase=True,items=[dict(item,reference=f'F-{i}') for i in range(count)],today=date.today(),selected_state='undated',query='',
                    summary=dict(overdue_amount=0,overdue_count=0,due_soon_amount=0,due_soon_count=0,open_amount=100*count))
    def test_excel_safe_text_and_unknown_due(self):
        ws=load_workbook(export_excel(self.context(2))).active
        self.assertEqual(ws['A10'].data_type,'s')
        self.assertEqual(ws['D10'].value,'Vadesi belirtilmemiş')
        self.assertEqual(ws.cell(ws.max_row,7).value,200)
        self.assertIn('Fatura',ws['A1'].value)
    def test_empty_and_long_pdf(self):
        for count in (0,80):
            self.assertTrue(export_pdf(self.context(count)).getvalue().startswith(b'%PDF'))
            self.assertTrue(export_excel(self.context(count)).getvalue().startswith(b'PK'))
if __name__=='__main__':unittest.main()
