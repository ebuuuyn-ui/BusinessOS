"""Read-only exports of the same filtered invoice ledger shown on screen."""
from datetime import date
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from xml.sax.saxutils import escape

ZERO = Decimal('0')
NOTE = 'Tutarlar TL. Pozitif bakiye borç, negatif bakiye alacaktır. Fatura ve cari hareketler esas alınır; bağlı siparişler ayrıca borçlandırılmaz.'


def statement_period(all_entries, start=None, end=None):
    opening = sum((e['debit'] - e['credit'] for e in all_entries if start and e['date'] < start), ZERO)
    entries = [e for e in all_entries if (not start or e['date'] >= start) and (not end or e['date'] <= end)]
    debit = sum((e['debit'] for e in entries), ZERO)
    credit = sum((e['credit'] for e in entries), ZERO)
    return dict(entries=entries, opening_balance=opening, period_debit=debit,
                period_credit=credit, closing_balance=opening + debit - credit)


def payment_details(entry):
    parts = [entry.get('payment_method') or '']
    for key, label in [('check_no', 'Çek'), ('check_bank', 'Banka'), ('check_due_date', 'Vade'),
                       ('check_status', 'Çek durumu'), ('card_installments', 'Taksit'),
                       ('card_owner_type', 'Kart sahibi'), ('card_customer_name', 'Müşteri')]:
        value = entry.get(key)
        if value:
            parts.append(f'{label}: {value:%d.%m.%Y}' if isinstance(value, date) else f'{label}: {value}')
    return ' · '.join(filter(None, parts))


def description(entry):
    return '\n'.join(filter(None, [entry.get('description'), payment_details(entry)]))


def period_label(start, end):
    return f"Tarih aralığı: {start or 'Başlangıçtan'} / {end or 'Son kayda kadar'}"


def item_values(item):
    return [item.product_name, getattr(item, 'variant', '') or '', getattr(item, 'detail_2', '') or '',
            getattr(item, 'detail_3', '') or '', item.quantity, item.unit, item.unit_price,
            'Dahil' if item.vat_included else 'Hariç', item.discount_rate, item.vat_rate,
            item.net_amount, item.vat_amount, item.total_amount, getattr(item, 'note', '') or '']


def export_xlsx(customer, period, start='', end='', detailed=False):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    book = Workbook()
    book.remove(book.active)

    def sheet(name, headers, rows, widths, numeric=()):
        ws = book.create_sheet(name)
        for text in [f'Cari Hesap Ekstresi - {customer.name}', f'Cari kodu: {customer.code or "-"} | {period_label(start, end)}', NOTE]:
            ws.append([text])
            ws.merge_cells(start_row=ws.max_row, start_column=1, end_row=ws.max_row, end_column=len(headers))
        ws.append(headers)
        for values in rows:
            ws.append([float(v) if isinstance(v, Decimal) else v for v in values])
        for row in ws:
            for cell in row:
                if cell.data_type == 'f':
                    cell.data_type = 's'
                cell.font = Font(name='Arial', size=10)
                cell.alignment = Alignment(vertical='top', wrap_text=True)
                if cell.row > 4 and cell.column in numeric:
                    cell.number_format = '#,##0.00;[Red]-#,##0.00'
                if isinstance(cell.value, date):
                    cell.number_format = 'dd.mm.yyyy'
                if cell.row > 4 and cell.row % 2:
                    cell.fill = PatternFill('solid', fgColor='F1F5FA')
        for cell in ws[4]:
            cell.fill = PatternFill('solid', fgColor='24456A')
            cell.font = Font(name='Arial', bold=True, color='FFFFFF', size=10)
        ws.cell(1, 1).font = Font(name='Arial', size=15, bold=True, color='14213D')
        ws.row_dimensions[1].height = 42
        ws.row_dimensions[2].height = 30
        ws.row_dimensions[3].height = 36
        ws.row_dimensions[4].height = 32
        for idx, width in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(idx)].width = width
        # Explicit heights avoid hidden wrapped descriptions in Excel's initial view.
        import math
        for row in ws.iter_rows(min_row=5):
            lines = max((sum(max(1, math.ceil(len(part) / max(8, widths[c.column-1] - 3))) for part in str(c.value or '').split('\n')) for c in row), default=1)
            ws.row_dimensions[row[0].row].height = min(409, max(26, lines * 14))
        ws.freeze_panes = 'C5'
        ws.auto_filter.ref = f'A4:{get_column_letter(len(headers))}{max(4,ws.max_row)}'
        ws.sheet_view.showGridLines = False
        ws.page_setup.orientation = 'landscape'
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.print_title_rows = '1:4'
        return ws

    entries = period['entries']
    rows = [[None, '', 'Dönem Başı Bakiye', None, None, period['opening_balance']]]
    rows += [[e['date'], e['reference'], description(e), e['debit'], e['credit'], e['balance']] for e in entries]
    ws = sheet('Ekstre', ['Tarih', 'Fatura / Referans No', 'Açıklama', 'Borç', 'Alacak', 'Bakiye'], rows, [16, 28, 70, 22, 22, 22], {4, 5, 6})
    ws.append([None, '', 'Dönem Toplamı / Son Bakiye', float(period['period_debit']), float(period['period_credit']), float(period['closing_balance'])])
    for cell in ws[ws.max_row]:
        cell.font = Font(name='Arial', bold=True, color='14213D')
        cell.fill = PatternFill('solid', fgColor='E3ECF7')
        if cell.column >= 4:
            cell.number_format = '#,##0.00;[Red]-#,##0.00'
    ws.row_dimensions[ws.max_row].height = 30
    if detailed:
        headers = ['Fatura No', 'Tarih', 'Bağlı Sipariş', 'Ürün / Hizmet', 'Ayrıntı 1', 'Ayrıntı 2', 'Ayrıntı 3', 'Adet', 'Birim', 'Birim Fiyat', 'Fiyat KDV', 'İskonto %', 'KDV %', 'KDV Hariç', 'KDV', 'KDV Dahil', 'Kalem Notu', 'Belge Notu']
        invoice_rows, order_rows, seen_orders = [], [], set()
        for e in entries:
            invoice = e.get('invoice')
            if not invoice:
                continue
            order = invoice.order
            for item in invoice.items:
                invoice_rows.append([invoice.invoice_no, invoice.invoice_date, order.order_no if order else ''] + item_values(item) + [invoice.notes or ''])
            if order and order.id not in seen_orders:
                seen_orders.add(order.id)
                for item in order.items:
                    order_rows.append(['Bilgi amaçlı', order.order_date, order.order_no] + item_values(item) + [order.notes or ''])
        widths = [26, 16, 24, 45, 25, 25, 25, 10, 12, 20, 12, 12, 12, 20, 20, 20, 40, 40]
        sheet('Fatura Kalemleri', headers, invoice_rows, widths, {10, 14, 15, 16})
        sheet('Bağlı Sipariş Kalemleri', ['Bilgi Amaçlı - Bakiyeye Eklenmez'] + headers[1:], order_rows, widths, {10, 14, 15, 16})
        sheet('Ödeme ve Tahsilat', ['Tarih', 'Referans', 'Açıklama', 'Ödeme Ayrıntıları', 'Borç', 'Alacak', 'Bakiye'],
              [[e['date'], e['reference'], e['description'], payment_details(e), e['debit'], e['credit'], e['balance']] for e in entries if e.get('transaction_id')],
              [16, 28, 60, 65, 22, 22, 22], {5, 6, 7})
    output = BytesIO()
    book.save(output)
    output.seek(0)
    return output


def export_pdf(customer, period, start='', end='', detailed=False):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak
    regular = Path('/System/Library/Fonts/Supplemental/Arial.ttf')
    bold = Path('/System/Library/Fonts/Supplemental/Arial Bold.ttf')
    if not regular.exists():
        regular = Path('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf')
        bold = Path('/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf')
    # Never silently render Turkish characters with an incompatible fallback.
    pdfmetrics.registerFont(TTFont('StatementFont', str(regular)))
    pdfmetrics.registerFont(TTFont('StatementBold', str(bold)))
    style = ParagraphStyle('cell', fontName='StatementFont', fontSize=8, leading=11, wordWrap='CJK')
    title = ParagraphStyle('title', parent=style, fontName='StatementBold', fontSize=16, leading=21, spaceAfter=8)
    heading = ParagraphStyle('heading', parent=style, fontName='StatementBold', fontSize=10, leading=14, spaceBefore=10, spaceAfter=6, keepWithNext=True)
    header = ParagraphStyle('header', parent=style, fontName='StatementBold', textColor=colors.white)
    right = ParagraphStyle('right', parent=style, alignment=2)
    para = lambda value, s=style: Paragraph(escape(str(value or '')).replace('\n', '<br/>'), s)
    money = lambda value: f'{value:,.2f}'.replace(',', 'X').replace('.', ',').replace('X', '.')
    output = BytesIO()
    doc = SimpleDocTemplate(output, pagesize=landscape(A4), leftMargin=28, rightMargin=28, topMargin=28, bottomMargin=30, title='Cari Hesap Ekstresi')
    width = landscape(A4)[0] - 56

    def table(headers, rows, proportions, amounts=()):
        data = [[para(h, header) for h in headers]]
        for row in rows:
            data.append([para(money(v), right) if idx in amounts and v is not None else para(v.strftime('%d.%m.%Y') if isinstance(v, date) else v) for idx, v in enumerate(row)])
        result = Table(data, colWidths=[width*p/sum(proportions) for p in proportions], repeatRows=1, splitByRow=1, splitInRow=1, hAlign='LEFT')
        result.setStyle(TableStyle([('BACKGROUND', (0,0),(-1,0),colors.HexColor('#24456A')), ('VALIGN',(0,0),(-1,-1),'TOP'), ('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.white,colors.HexColor('#F1F5FA')]), ('LINEBELOW',(0,0),(-1,-1),.3,colors.HexColor('#D6DFEA')), ('LEFTPADDING',(0,0),(-1,-1),6), ('RIGHTPADDING',(0,0),(-1,-1),6), ('TOPPADDING',(0,0),(-1,-1),6), ('BOTTOMPADDING',(0,0),(-1,-1),6)]))
        return result

    story = [para('Ayrıntılı Cari Hesap Ekstresi' if detailed else 'Cari Hesap Ekstresi', title), para(customer.name, heading), para(f'Cari kodu: {customer.code or "-"} | {period_label(start,end)}'), para(NOTE), Spacer(1,10)]
    story.append(table(['Dönem Başı Bakiye','Dönem Borç','Dönem Alacak','Dönem Sonu Bakiye'], [[period[k] for k in ['opening_balance','period_debit','period_credit','closing_balance']]], [1]*4, range(4)))
    story.append(Spacer(1,12))
    rows = [[e['date'],e['reference'],description(e),e['debit'],e['credit'],e['balance']] for e in period['entries']]
    rows.append(['','','Dönem Toplamı / Son Bakiye',period['period_debit'],period['period_credit'],period['closing_balance']])
    story.append(table(['Tarih','Fatura / Referans No','Açıklama','Borç','Alacak','Bakiye'],rows,[10,17,34,13,13,13],{3,4,5}))
    if not period['entries']:
        story.append(para('Bu tarih aralığında cari hesap hareketi bulunmuyor.'))
    if detailed and period['entries']:
        story.extend([PageBreak(), para('Hareket Ayrıntıları', title)])
        for e in period['entries']:
            story.append(para(f"{e['date']:%d.%m.%Y} | {e['reference']} | {e['description'] or ''}", heading))
            story.append(para(f"Borç: {money(e['debit'])} TL | Alacak: {money(e['credit'])} TL | Bakiye: {money(e['balance'])} TL"))
            invoice = e.get('invoice')
            if not invoice:
                story.append(para(payment_details(e)))
                continue
            if invoice.notes:
                story.append(para(f'Fatura notu: {invoice.notes}'))
            if invoice.due_date:
                story.append(para(f'Fatura vadesi: {invoice.due_date:%d.%m.%Y}'))
            documents = [('Fatura Kalemleri', invoice.items)]
            if invoice.order:
                o = invoice.order
                documents.append((f'Bağlı Sipariş {o.order_no} - {o.status} (bilgi amaçlı, bakiyeye eklenmez)', o.items))
            for label, items in documents:
                story.append(para(label, heading))
                values = []
                for i in items:
                    details = '\n'.join(filter(None,[i.product_name,getattr(i,'variant',''),getattr(i,'detail_2',''),getattr(i,'detail_3',''),getattr(i,'note','')]))
                    values.append([details,f'{i.quantity} {i.unit}',f'{money(i.unit_price)}\nKDV {"dahil" if i.vat_included else "hariç"}',f'%{money(i.discount_rate)}',f'%{money(i.vat_rate)}',i.net_amount,i.vat_amount,i.total_amount])
                story.append(table(['Ürün / Hizmet ve Ayrıntılar','Adet','Birim Fiyat','İskonto','KDV %','KDV Hariç','KDV','KDV Dahil'],values,[32,8,12,7,7,12,10,12],{5,6,7}))
            if invoice.order and invoice.order.notes:
                story.append(para(f'Sipariş notu: {invoice.order.notes}'))

    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont('StatementFont', 8)
        canvas.drawString(28, 15, f'Business OS | {date.today():%d.%m.%Y}')
        canvas.drawRightString(width+28, 15, f'Sayfa {doc.page}')
        canvas.restoreState()
    doc.build(story,onFirstPage=footer,onLaterPages=footer)
    output.seek(0)
    return output
