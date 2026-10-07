import io,json,unittest
from decimal import Decimal
from openpyxl import load_workbook
from pypdf import PdfReader
from test_price_quotes import QuoteTests
from test_web_auth import BASE
import app as m
from packing_lists import defaults,validate,pdf,excel

class PackingTests(unittest.TestCase):
    setUp=QuoteTests.setUp
    tearDown=QuoteTests.tearDown
    sample=QuoteTests.sample
    post_quote=QuoteTests.post
    def prepare(self):
        d=self.sample();self.post_quote(d);q=m.PriceQuote.query.one();source=json.loads(q.payload);packing=defaults(q,source)
        packing['rows'][0].update(name='Office chair',origin='Türkiye',hs='940139000000',cartons='2',gross='46',net='40',length='100',width='50',height='80',marks='1–2')
        return q,source,packing
    def post(self,q,data,version='1'):
        return self.client.post(f'/fiyat-teklifleri/{q.id}/packing-list',base_url=BASE,headers={'Origin':BASE},data={'version':version,'payload':json.dumps(data)})
    def test_save_export_preserve_and_no_accounting(self):
        q,source,data=self.prepare();order_count=m.Order.query.count();stock_count=m.StockMovement.query.count()
        self.assertEqual(self.post(q,data).status_code,302)
        m.db.session.expire_all();q=m.db.session.get(m.PriceQuote,q.id);saved=json.loads(q.payload)
        self.assertEqual(saved['totals'],source['totals']);self.assertEqual(saved['packing_list']['totals']['cbm'],'0.800')
        self.assertEqual(saved['packing_list']['totals']['quantity'],4);self.assertEqual(saved['packing_list']['totals']['cartons'],2)
        for kind in ('pdf','xlsx'):
            r=self.client.get(f'/fiyat-teklifleri/{q.id}/packing-list/{kind}',base_url=BASE);self.assertEqual(r.status_code,200)
            if kind=='pdf':
                text=' '.join(p.extract_text() for p in PdfReader(io.BytesIO(r.data)).pages);self.assertIn('PACKING LIST',text);self.assertIn('Office chair',text);self.assertNotIn('5100',text)
            else:
                ws=load_workbook(io.BytesIO(r.data)).active;self.assertEqual(ws['F11'].value,4);self.assertEqual(ws['G11'].value,2);self.assertEqual(ws['J11'].value,.8)
        self.assertIn('başka bir oturumda',self.post(q,data,'1').text)
        edit=dict(source);self.assertEqual(self.post_quote(edit,q.id,str(q.version)).status_code,302)
        m.db.session.expire_all();self.assertIn('packing_list',json.loads(m.db.session.get(m.PriceQuote,q.id).payload))
        self.assertEqual(m.Order.query.count(),order_count);self.assertEqual(m.StockMovement.query.count(),stock_count);self.assertEqual(m.Invoice.query.count(),0)
    def test_split_partial_manual_volume_and_validation(self):
        q,source,d=self.prepare();first=d['rows'][0];first.update(quantity=1,cartons=1,length='',width='',height='',cbm='0.25')
        d['rows'].append({**first,'quantity':3,'cartons':2,'cbm':'0.5'})
        self.assertEqual(validate(d,source)['totals']['cbm'],'0.750')
        for field,value in [('quantity',4),('gross',39),('net','NaN'),('hs','123'),('cartons',1.5),('length',1)]:
            bad=json.loads(json.dumps(d));bad['rows'][1][field]=value
            with self.assertRaises(ValueError):validate(bad,source)
        clean=validate(d,source);clean['rows'][0]['name']='=1+1';ws=load_workbook(excel(clean,q.number)).active
        self.assertEqual(ws['C11'].data_type,'s')
    def test_access_and_missing_exports(self):
        q,source,d=self.prepare();path=f'/fiyat-teklifleri/{q.id}/packing-list'
        self.assertEqual(self.app.test_client().get(path,base_url=BASE).status_code,302)
        self.assertEqual(self.client.get(path+'/pdf',base_url=BASE).status_code,404)
        self.assertEqual(self.client.post(path,base_url=BASE,data={}).status_code,403)
        self.assertEqual(self.client.get(path,base_url=BASE).status_code,200)
