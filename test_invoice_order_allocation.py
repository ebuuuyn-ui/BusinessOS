import unittest
from unittest.mock import patch
from decimal import Decimal
import test_uyumsoft_inbox as fixtures
import app as m
from test_web_auth import BASE
from invoice_order_allocation import summaries, remaining, filter_orders
from invoice_export import build_invoice_register
from openpyxl import load_workbook

class AllocationTests(unittest.TestCase):
    setUp=fixtures.InboxTests.setUp
    tearDown=fixtures.InboxTests.tearDown
    fixture=fixtures.InboxTests.fixture
    connect=fixtures.InboxTests.connect
    make_invoice=fixtures.InboxTests.make_invoice
    post=fixtures.InboxTests.post
    mock=fixtures.InboxTests.mock
    preview=fixtures.InboxTests.preview
    def orders(self,second_quantity=1):
        self.fixture()
        a=m.db.session.get(m.Order,self.order_id)
        a.items=[m.OrderItem(product=self.product,product_name=self.product.name,quantity=1,unit='Adet')]
        b=m.Order(order_no='SA-TEST-2',order_type='Satın Alma',customer_id=1339)
        b.items=[m.OrderItem(product=self.product,product_name=self.product.name,quantity=second_quantity,unit='Adet')]
        m.db.session.add(b);m.db.session.commit();self.orders_pair=(a,b)
        self.form.update({'allocation_item_0[]':[str(a.items[0].id),str(b.items[0].id)],'allocation_qty_0[]':['1','1']})
    def test_one_invoice_two_orders_one_stock_effect_and_pdf(self):
        self.orders()
        with patch('uyumsoft_inbox.Client') as c:
            self.mock(c);self.preview();self.assertEqual(self.post(self.path,self.form).status_code,302)
            self.assertEqual(self.post(self.path,self.form).status_code,302)
        inv=m.Invoice.query.one();self.assertIsNone(inv.order_id)
        self.assertEqual(inv.total_amount,Decimal('198'))
        self.assertEqual(m.StockMovement.query.count(),1)
        self.assertEqual(m.StockMovement.query.one().quantity,2)
        self.assertEqual(m.InvoiceOrderAllocation.query.count(),2)
        self.assertEqual({o.id for o in inv.linked_orders},{o.id for o in self.orders_pair})
        for order in self.orders_pair:
            self.assertEqual(summaries([order])[order.id]['state'],'Faturalandı')
            page=self.client.get(f'/siparisler/{order.id}',base_url=BASE)
            self.assertIn('INB2026000000001',page.text)
            self.assertIn('Fatura – sipariş dağıtımı',page.text)
            self.assertIn('bağlı fatura',self.client.get(f'/siparisler/{order.id}/sil',base_url=BASE).text)
            self.assertEqual(self.post(f'/siparisler/{order.id}/duzenle',{}).status_code,302)
        listed=self.client.get('/siparisler?invoice_status=Faturalandı',base_url=BASE)
        self.assertEqual(listed.status_code,200);self.assertIn('SA-TEST-2',listed.text)
        detail=self.client.get(f'/faturalar/{inv.id}',base_url=BASE)
        self.assertIn('SA-TEST-2',detail.text);self.assertIn('SA-TEST',detail.text)
        sheet=load_workbook(build_invoice_register([inv])).active
        self.assertIn('SA-TEST-2',sheet.cell(2,9).value)
    def test_partial_quantity_remains_open_and_filters_agree(self):
        self.orders(3)
        with patch('uyumsoft_inbox.Client') as c:
            self.mock(c);self.preview();self.post(self.path,self.form)
        a,b=self.orders_pair
        self.assertEqual(remaining(b.items)[b.items[0].id],2)
        self.assertEqual(summaries([b])[b.id]['state'],'Kısmen Faturalandı')
        self.assertEqual({o.id for o in filter_orders(m.Order.query,'Faturalandı')},{a.id})
        self.assertEqual({o.id for o in filter_orders(m.Order.query,'Kısmen Faturalandı')},{b.id})
        self.assertEqual({o.id for o in filter_orders(m.Order.query,'Fatura Bekliyor')},{b.id})
    def test_over_allocation_and_wrong_product_rejected(self):
        self.orders()
        with patch('uyumsoft_inbox.Client') as c:
            self.mock(c);self.preview()
            wrong=dict(self.form,**{'allocation_qty_0[]':['2','1']})
            self.assertIn('aşılıyor',self.post(self.path,wrong).text)
            wrong=dict(self.form,**{'allocation_item_0[]':[self.form['allocation_item_0[]'][0]],'allocation_qty_0[]':['1']})
            self.assertIn('eşit olmalıdır',self.post(self.path,wrong).text)
            second=self.orders_pair[1].items[0];second.product_id=None;m.db.session.commit();self.preview()
            self.assertIn('stok kartı',self.post(self.path,self.form).text)
        self.assertEqual(m.Invoice.query.count(),0);self.assertEqual(m.StockMovement.query.count(),0)
    def test_stale_order_and_wrong_supplier_cannot_be_booked(self):
        self.orders()
        with patch('uyumsoft_inbox.Client') as c:
            self.mock(c);self.preview()
            self.orders_pair[1].items[0].quantity=3;m.db.session.commit()
            self.assertIn('değişmiş',self.post(self.path,self.form).text)
            self.orders_pair[1].customer=m.Customer(name='Another supplier',tax_number='9999999999');m.db.session.commit();self.preview()
            self.assertIn('geçerli sipariş',self.post(self.path,self.form).text)
        self.assertEqual(m.Invoice.query.count(),0)
    def test_order_filter_requires_all_lines_and_rejects_unknown_orders(self):
        self.orders()
        with patch('uyumsoft_inbox.Client') as c:
            self.mock(c)
            result=self.client.get(self.path+'?order_refs=UNKNOWN',base_url=BASE)
            self.assertIn('kontrol edin',result.text)
            self.path+='?order_refs=SA-TEST,SA-TEST-2';self.preview()
            form=dict(self.form,**{'allocation_item_0[]':[''],'allocation_qty_0[]':['']})
            self.assertIn('dağıtın',self.post(self.path,form).text)
        self.assertEqual(m.Invoice.query.count(),0)
    def test_legacy_invoice_keeps_old_filter_behavior(self):
        self.orders()
        a,b=self.orders_pair
        inv=m.Invoice(invoice_no='MANUAL',invoice_type='Satın Alma',customer_id=1339,order_id=a.id)
        m.db.session.add(inv);m.db.session.commit()
        self.assertEqual({o.id for o in filter_orders(m.Order.query,'Faturalandı')},{a.id})
        self.assertEqual({o.id for o in filter_orders(m.Order.query,'Fatura Bekliyor')},{b.id})
        with patch('uyumsoft_inbox.Client') as c:
            self.mock(c);self.preview()
            self.assertIn('geçerli sipariş',self.post(self.path,self.form).text)

    def test_existing_database_without_allocation_table(self):
        self.orders()
        m.InvoiceOrderAllocation.__table__.drop(m.db.engine)
        self.assertEqual(filter_orders(m.Order.query,'Kısmen Faturalandı').count(),0)
        self.assertEqual(filter_orders(m.Order.query,'Fatura Bekliyor').count(),2)
        self.assertEqual(filter_orders(m.Order.query,'Faturalandı').count(),0)
        for order in self.orders_pair:
            self.assertEqual(summaries([order])[order.id]['state'],'Fatura Bekliyor')
